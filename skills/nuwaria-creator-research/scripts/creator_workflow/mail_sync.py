"""Reviewable Zoho-to-Notion status updates. This module never sends email."""
import argparse, json, time
from pathlib import Path
from datetime import datetime, timezone, timedelta
from outreach import Run, read, save, flat, rt, para, text, EMAIL, digest, now, status_names

def local_date(m):return datetime.fromtimestamp(int(m['receivedTime'])/1000,timezone(timedelta(hours=8))).date().isoformat()
def typed(meta,value):
    typ=meta['type']
    if typ in ('title','rich_text'):return {typ:rt(value)}
    if typ=='select':
        if value is None:return {'select':None}
        assert value in {x['name'] for x in meta['select']['options']}
        return {'select':{'name':value}}
    if typ=='date':return {'date':{'start':value} if value else None}
    if typ=='checkbox':return {'checkbox':bool(value)}
    if typ=='multi_select':
        names=value if isinstance(value,(list,set,tuple)) else [value] if value else []
        assert set(names)<={x['name'] for x in meta['multi_select']['options']}
        return {'multi_select':[{'name':n} for n in sorted(set(names))]}
    if typ in ('email','phone_number','url'):return {typ:value or None}
    raise ValueError('Unsupported property '+typ)

class MailSync:
    def __init__(self,project,folder):self.r=Run(project);self.o=Path(folder).resolve()
    def collect(self):
        from bs4 import BeautifulSoup
        day=read(self.o/'collection_scope.json')['date'];datetime.fromisoformat(day)
        save(self.o/'notion_before.json',self.r.all_pages())
        save(self.r.o/'schema_live.json',self.r.notion('GET','data_sources/'+self.r.ds))
        save(self.o/'tracking_schema.json',self.r.notion('GET','data_sources/'+self.r.c['notion_tracking_data_source']))
        rows=[];cursor=None
        while True:
            body={'page_size':100}
            if cursor:body['start_cursor']=cursor
            result=self.r.notion('POST','data_sources/'+self.r.c['notion_tracking_data_source']+'/query',json=body);rows.extend(result['results'])
            if not result.get('has_more'):break
            cursor=result['next_cursor']
        save(self.o/'tracking_pages.json',{'results':rows,'has_more':False})
        folders=self.r.zoho('GET','/folders')['data'];save(self.o/'folders.json',folders)
        sent=self.r.sent();save(self.o/'sent_all.json',sent)
        selected={str(m['messageId']):{'meta':m,'direction':'sent','folder_id':self.r.sent_folder} for m in sent if local_date(m)==day}
        scanned=[]
        for f in folders:
            if f.get('folderType','').lower() not in ('inbox','archive','snoozed') and f.get('folderName','').lower() not in ('inbox','archive','snoozed'):continue
            count=0
            for start in range(1,20001,100):
                batch=self.r.zoho('GET','/messages/view',params={'folderId':f['folderId'],'limit':100,'start':start})['data'];count+=len(batch)
                for m in batch:
                    if local_date(m)!=day:continue
                    system=any(t in (str(m.get('fromAddress',''))+' '+m.get('subject','')).lower() for t in ['mailer-daemon','postmaster','undeliver','delivery status'])
                    selected[str(m['messageId'])]={'meta':m,'direction':'system' if system else 'inbox','folder_id':str(f['folderId'])}
                if len(batch)<100:break
            else:raise RuntimeError('Incomplete folder pagination')
            scanned.append({'folder':f['folderName'],'messages':count})
        messages=[]
        for i,m in enumerate(sorted(selected.values(),key=lambda x:int(x['meta']['receivedTime'])),1):
            raw=self.r.zoho('GET',f"/folders/{m['folder_id']}/messages/{m['meta']['messageId']}/content")['data']['content']
            m.update(raw_content=raw,full_text=BeautifulSoup(raw,'html.parser').get_text('\n',strip=True) if '<' in raw else raw)
            messages.append(m);save(self.o/'message_contents.json',messages)
            if i%10==0:print('MAIL_READ',i,'/',len(selected),flush=True)
        save(self.o/'collection_report.json',{'date':day,'timezone':'Asia/Shanghai','collected_at':now(),'folders':scanned,'sent_scanned':len(sent),'messages':len(messages)})
        print('COLLECTED',len(messages),'messages for',day,flush=True)
    def attachments(self):
        destination=self.o/'attachments';destination.mkdir(exist_ok=True)
        manifest=[]
        selected=read(self.o/'attachment_scope.json') if (self.o/'attachment_scope.json').exists() else None
        for m in read(self.o/'message_contents.json'):
            if selected is not None and str(m['meta']['messageId']) not in selected:continue
            if selected is not None and str(m['meta'].get('hasInline','')).lower()=='true':
                from email import policy
                from email.parser import Parser
                mid=str(m['meta']['messageId'])
                original=self.r.zoho('GET',f'/messages/{mid}/originalmessage')['data']['content']
                mime=Parser(policy=policy.default).parsestr(original)
                for i,part in enumerate(mime.walk()):
                    if part.get_content_type() not in ('image/png','image/jpeg','image/webp'):continue
                    data=part.get_payload(decode=True)
                    if not data or len(data)>30000000:continue
                    path=destination/(mid+'_inline_'+str(i)+'.'+part.get_content_subtype());path.write_bytes(data)
                    manifest.append({'message_id':mid,'filename':path.name,'path':str(path),'inline':True})
                    save(self.o/'attachment_manifest.json',manifest);print('INLINE_SAVED',path.name,flush=True)
            if str(m['meta'].get('hasAttachment','0'))!='1':continue
            base=f"/folders/{m['folder_id']}/messages/{m['meta']['messageId']}"
            info=self.r.zoho('GET',base+'/attachmentinfo')['data']
            for a in info.get('attachments',[]):
                filename=Path(a['attachmentName']).name
                if Path(filename).suffix.lower() not in ('.pdf','.png','.jpg','.jpeg','.webp'):continue
                assert int(a.get('attachmentSize',0))<30000000
                path=destination/(str(m['meta']['messageId'])+'_'+filename)
                response=self.r.http.get(self.r.zbase+base+'/attachments/'+str(a['attachmentId']),headers=self.r.zh,timeout=45)
                response.raise_for_status();path.write_bytes(response.content)
                manifest.append({'message_id':str(m['meta']['messageId']),'filename':filename,'path':str(path),'attachment_id':str(a['attachmentId'])})
                save(self.o/'attachment_manifest.json',manifest)
                print('ATTACHMENT_SAVED',filename,len(response.content),flush=True)
    def reply_audit(self):
        """Refresh recent mail with overlap; retain historical context and daily full reconciliation."""
        from bs4 import BeautifulSoup
        from mail_cache import scan_folder
        pages=self.r.all_pages();save(self.o/'notion_before.json',pages)
        known={e.lower() for p in pages for e in EMAIL.findall(str(flat(p).get('公开联系方式','')))}
        folders=self.r.zoho('GET','/folders')['data'];save(self.o/'folders.json',folders)
        sent=self.r.sent();save(self.o/'sent_all.json',sent)
        incoming={};scanned=[]
        for f in folders:
            kind=f.get('folderType','').lower();name=f.get('folderName','').lower()
            if kind in ('sent','draft','drafts','trash','spam','outbox') or name in ('sent','drafts','trash','spam','outbox'):continue
            rows=scan_folder(self.r,f['folderId'])
            scanned.append({'name':f['folderName'],'count':len(rows)})
            for m in rows:
                sender=str(m.get('fromAddress',''));subject=m.get('subject','')
                if any(k in (sender+' '+subject).lower() for k in ['mailer-daemon','postmaster','undeliver','delivery status','delivery failure']):continue
                if any(term in subject.lower() for term in ('nuwaria','collaboration','partnership','ugc')) or known&{e.lower() for e in EMAIL.findall(sender)}:
                    incoming[str(m['messageId'])]={'meta':m,'direction':'inbox','folder_id':str(f['folderId'])}
        addresses={e.lower() for m in incoming.values() for e in EMAIL.findall(str(m['meta'].get('fromAddress','')))}
        subjects={m['meta'].get('subject','').lower().replace('re:','').replace('fwd:','').strip() for m in incoming.values()}
        latest={}
        for m in sorted(sent,key=lambda m:int(m['receivedTime'])):
            to={e.lower() for e in EMAIL.findall(str(m.get('toAddress','')))}
            subject=m.get('subject','').lower().replace('re:','').replace('fwd:','').strip()
            if to&addresses or subject in subjects:
                for e in to:latest[e]=m
        selected=dict(incoming)
        for m in latest.values():selected[str(m['messageId'])]={'meta':m,'direction':'sent','folder_id':self.r.sent_folder}
        cache_path=self.r.o/'mail_cache'/'bodies.json'
        cache=read(cache_path) if cache_path.exists() else {}
        if not cache:
            for path in self.o.parent.glob('*/message_contents.json'):
                for m in read(path):cache[str(m['meta']['messageId'])]=m
        messages=[];downloaded=0;cache_hits=0
        for i,m in enumerate(sorted(selected.values(),key=lambda m:int(m['meta']['receivedTime'])),1):
            mid=str(m['meta']['messageId'])
            if mid in cache and 'full_text' in cache[mid]:
                m.update(raw_content=cache[mid].get('raw_content',''),full_text=cache[mid]['full_text'])
                cache_hits+=1
            else:
                raw=self.r.zoho('GET',f"/folders/{m['folder_id']}/messages/{mid}/content")['data']['content']
                m.update(raw_content=raw,full_text=BeautifulSoup(raw,'html.parser').get_text('\n',strip=True) if '<' in raw else raw)
                downloaded+=1;print('BODY_FETCH',downloaded,flush=True)
            messages.append(m);cache[mid]=m
        save(cache_path,cache);save(self.o/'message_contents.json',messages)
        scans=getattr(self.r,'mail_scan_reports',[])
        new_ids={mid for scan in scans for mid in scan['new_message_ids']}
        save(self.o/'new_messages.json',[m for m in messages if str(m['meta']['messageId']) in new_ids])
        save(self.o/'collection_report.json',{'checked_at':now(),'folders':scanned,'incoming_messages':len(incoming),'latest_outgoing_messages':len(selected)-len(incoming),'sent_metadata':len(sent),'body_cache_hits':cache_hits,'bodies_downloaded':downloaded,'metadata_pages_fetched':sum(s['pages_fetched'] for s in scans),'folder_scans':scans,'excluded_folders':'Trash, Spam, Drafts, Sent, Outbox; Sent scanned separately','writes_or_sends':0})
        report=read(self.o/'collection_report.json')
        report['folder_scans']=[{k:v for k,v in s.items() if k!='new_message_ids'} for s in scans]
        print(json.dumps(report,ensure_ascii=False),flush=True)
    def reply_report(self):
        from collections import Counter
        from openpyxl import Workbook,load_workbook
        from openpyxl.styles import Font,PatternFill,Alignment
        decisions=read(self.o/'reply_triage.json');pages=read(self.o/'notion_before.json')
        messages=read(self.o/'message_contents.json');sent=read(self.o/'sent_all.json')
        zone=timezone(timedelta(hours=8))
        stamp=lambda m:datetime.fromtimestamp(int(m['receivedTime'])/1000,zone).strftime('%Y-%m-%d %H:%M')
        wb=Workbook();todo=wb.active;todo.title='需要我方处理';allrows=wb.create_sheet('全部有回复达人');waiting=wb.create_sheet('原待跟进123人拆分')
        headers=['达人','账号','处理类型','下一步','判断依据','最新来信时间','最新发信时间','往来邮箱','Notion','来信ID','发信ID']
        todo.append(headers);allrows.append(headers);waiting.append(['账号','拆分结果','下一步','Notion'])
        triaged={};audit=[]
        for d in decisions:
            match=d.get('record_email',d['email']).lower()
            matches=[p for p in pages if match in {e.lower() for e in EMAIL.findall(str(flat(p).get('公开联系方式','')))}]
            assert len(matches)==1,(d['name'],len(matches))
            page=matches[0];f=flat(page)
            incoming=[m for m in messages if m['direction']=='inbox' and d['email'].lower() in {e.lower() for e in EMAIL.findall(str(m['meta'].get('fromAddress','')))}]
            assert incoming,d['name']
            last_in=max(incoming,key=lambda m:int(m['meta']['receivedTime']))['meta']
            emails={d['email'].lower()}|{e.lower() for e in EMAIL.findall(str(f.get('公开联系方式','')))}
            outgoing=[m for m in sent if emails&{e.lower() for e in EMAIL.findall(str(m.get('toAddress','')))}]
            last_out=max(outgoing,key=lambda m:int(m['receivedTime'])) if outgoing else None
            row=[d['name'],f['账号'],d['category'],d['next'],d['reason'],stamp(last_in),stamp(last_out) if last_out else '',d['email'],page['url'],str(last_in['messageId']),str(last_out['messageId']) if last_out else '']
            allrows.append(row)
            if d['category'] in ('优先回复','先决定再回复','寄样执行'):todo.append(row)
            triaged[page['id']]=d
            audit.append(d|{'page_id':page['id'],'handle':f['账号'],'latest_incoming_id':str(last_in['messageId']),'latest_outgoing_id':str(last_out['messageId']) if last_out else None})
        buckets=Counter()
        for page in pages:
            f=flat(page)
            if (f.get('联系状态') or {}).get('name')!='待跟进':continue
            d=triaged.get(page['id']);category=d['category'] if d else '未见实质来信（含退信待处理）'
            buckets[category]+=1
            waiting.append([f['账号'],category,d['next'] if d else f.get('下一步',''),page['url']])
        counts=Counter(d['category'] for d in decisions)
        info=wb.create_sheet('核验口径');info.append(['项目','说明'])
        info.append(['核验时间',read(self.o/'collection_report.json')['checked_at']])
        info.append(['范围','完整收件及自定义通知文件夹、归档、稍后处理；完整已发送。排除垃圾邮件、已删除、草稿，排除自动回复和表情通知作为待答问题。'])
        info.append(['判断','按人及邮箱核对，不因最新邮件来自对方就一律判为需回复，也不按同名邮件主题跨达人合并。寄样任务独立于回复；已报价不代表已批准。'])
        info.append(['分类数量',json.dumps(dict(counts),ensure_ascii=False)])
        info.append(['原待跟进拆分',json.dumps(dict(buckets),ensure_ascii=False)])
        for ws in wb:
            ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
            for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='405D51')
            for col in 'ABCDEFGHIJK':ws.column_dimensions[col].width=25
            for row in ws.iter_rows(min_row=2):
                ws.row_dimensions[row[0].row].height=65
                for c in row:
                    c.alignment=Alignment(wrap_text=True,vertical='top')
                    if isinstance(c.value,str):c.data_type='s'
        for ws in (todo,allrows):ws.column_dimensions['D'].width=70;ws.column_dimensions['E'].width=70
        waiting.column_dimensions['C'].width=80;info.column_dimensions['B'].width=110
        path=self.o/'NUWARIA_需要我回复与处理_2026-09-16.xlsx';wb.save(path)
        check=load_workbook(path,read_only=True);assert check['需要我方处理'].max_row==15;assert check['全部有回复达人'].max_row==40;check.close()
        save(self.o/'reply_audit_result.json',{'counts':dict(counts),'waiting_status_breakdown':dict(buckets),'reviewed':audit,'report':str(path)})
        print(json.dumps({'counts':dict(counts),'waiting_status_breakdown':dict(buckets),'report':str(path)},ensure_ascii=False),flush=True)

    def prepare_status_correction(self):
        decisions=read(self.o/'reply_triage.json');pages=read(self.o/'notion_before.json')
        messages=read(self.o/'message_contents.json');sent=read(self.o/'sent_all.json');reviewed={}
        for d in decisions:
            email=d.get('record_email',d['email']).lower()
            matches=[p for p in pages if email in {e.lower() for e in EMAIL.findall(str(flat(p).get('公开联系方式','')))}]
            assert len(matches)==1;reviewed[matches[0]['id']]=d
        plan=[]
        for page in pages:
            f=flat(page);old=(f.get('联系状态') or {}).get('name');d=reviewed.get(page['id']);values={};evidence=[]
            if d:
                category=d['category']
                status={'优先回复':'待跟进','先决定再回复':'待跟进','等待对方':'待二次回复','已回应，内部评估或留档':'已报价','已结束或暂缓':'已暂停','寄样执行':'待寄样'}[category]
                if d.get('needs_our_reply'):status='待跟进'
                incoming=[m for m in messages if m['direction']=='inbox' and d['email'].lower() in {e.lower() for e in EMAIL.findall(str(m['meta'].get('fromAddress','')))}]
                substantive=[m for m in incoming if 'reacted' not in m['full_text'][:200].lower()]
                last=max(substantive or incoming,key=lambda m:int(m['meta']['receivedTime']))
                emails={d['email'].lower()}|{e.lower() for e in EMAIL.findall(str(f.get('公开联系方式','')))}
                out=[m for m in sent if emails&{e.lower() for e in EMAIL.findall(str(m.get('toAddress','')))}]
                evidence=[str(last['meta']['messageId'])]
                if out:evidence.append(str(max(out,key=lambda m:int(m['receivedTime']))['messageId']))
                note='状态口径修正：'+d['reason']+' 下一步：'+d['next']
                values={'联系状态':{'select':{'name':status}},'下一步':{'rich_text':rt(d['next'])},'建联备注':{'rich_text':rt(f.get('建联备注','')+'\n\n'+note)}}
                # Reaction-only notifications do not establish a new substantive reply date.
                if substantive and not d.get('reaction_only',False):values['回复日期']={'date':{'start':local_date(last['meta'])}}
                if status=='待跟进':values['是否暂停推进']={'checkbox':False}
            elif old=='待跟进':
                # Previously mislabelled cold outreach; no reviewed substantive reply exists.
                values={'联系状态':{'select':{'name':'已发信'}}}
            if not values:continue
            changes={}
            for k,v in values.items():
                try:self.r.verify_properties(page,{k:v})
                except AssertionError:changes[k]=v
            if changes:plan.append({'table':'main','handle':f['账号'],'page_id':page['id'],'properties':changes,'before':{k:page['properties'][k] for k in changes},'message_ids':evidence})
        save(self.o/'update_plan.json',plan)
        from collections import Counter
        expected={p['id']:(flat(p).get('联系状态') or {}).get('name') for p in pages}
        for x in plan:
            if '联系状态' in x['properties']:expected[x['page_id']]=x['properties']['联系状态']['select']['name']
        save(self.o/'status_correction_preview.json',{'updates':len(plan),'expected_counts':dict(Counter(expected.values())),'needs_our_reply':[flat(p)['账号'] for p in pages if expected[p['id']]=='待跟进']})
        print(json.dumps(read(self.o/'status_correction_preview.json'),ensure_ascii=False),flush=True)

    def prepare(self):
        pages=read(self.o/'notion_before.json');tracking=read(self.o/'tracking_pages.json');assert not tracking['has_more']
        schema=read(self.r.o/'schema_live.json')['properties'];ts=read(self.o/'tracking_schema.json')['properties']
        messages={str(x['meta']['messageId']):x for x in read(self.o/'message_contents.json')};plan=[]
        for d in read(self.o/'decisions.json'):
            match_email=d.get('record_email',d['email'])
            matches=[p for p in pages if match_email.lower() in {e.lower() for e in EMAIL.findall(str(flat(p).get('公开联系方式','')))}]
            assert len(matches)==1,(d['email'],len(matches))
            page=matches[0];f=flat(page);ms=[messages[mid] for mid in d['messages']]
            for m in ms:
                meta=m['meta'];field='fromAddress' if m['direction']=='inbox' else 'toAddress'
                evidence=m.get('full_text','') if m['direction']=='system' else str(meta.get(field,''))
                evidence_emails={e.lower() for e in EMAIL.findall(evidence)}
                valid=d['email'].lower() in evidence_emails
                if not valid and m['direction']=='inbox':
                    alias_path=self.r.p/'contact_email_aliases.json'
                    aliases=read(alias_path) if alias_path.exists() else []
                    valid=any(a['handle']==f.get('账号') and a['current_email']==d['email'].lower() and set(a['previous_emails'])&evidence_emails for a in aliases)
                assert valid,'Message identity does not match current contact or verified historical alias'
            incoming=[m for m in ms if m['direction']=='inbox']
            if d['status']=='已暂停':
                assert incoming or f.get('回复日期'),'Paused collaboration requires creator-reply evidence; delivery failures are not replies'
            if d['status']=='待跟进':
                assert d.get('needs_our_reply') is True,'待跟进 requires an explicit reviewed need for our reply'
                assert incoming or f.get('回复日期'),'Cold outreach cannot be 待跟进'
            sources='；'.join({'sent':'发出','inbox':'收到','system':'系统通知'}[m['direction']]+' Zoho '+str(m['meta']['messageId']) for m in sorted(ms,key=lambda m:int(m['meta']['receivedTime'])))
            note=d['summary'].strip()
            if d['next'].strip() and d['next'].strip() not in note:note+=' '+d['next'].strip()
            values={'联系状态':d['status'],'是否暂停推进':d['paused'],'建联备注':note,'下一步':d['next'],'合作方式':d['cooperation'],**d.get('extra',{})}
            if schema['联系状态']['type']=='multi_select':
                tags={values['联系状态']} if isinstance(values['联系状态'],str) else set(values['联系状态'])
                tags|=status_names(f.get('联系状态'))&{'已报价'}
                if d.get('quote_received') is True:
                    assert d.get('quote_evidence'),'Reviewed quote evidence required'
                    tags.add('已报价')
                values['联系状态']=tags
            if incoming:
                latest=max(incoming,key=lambda m:int(m['meta']['receivedTime']));values['回复日期']=local_date(latest['meta'])
            if d.get('reason'):values['降级/淘汰原因']=d['reason']
            changes={k:typed(schema[k],v) for k,v in values.items()}
            marker='邮件同步 '+','.join(sorted(d['messages']))
            history=[para(marker),para('最新结论：'+d['summary']),para('下一步：'+d['next'])]
            if f.get('建联备注'):history.append(para('同步前备注（历史，不代表当前进度）：'+f['建联备注']))
            plan.append({'table':'main','handle':f['账号'],'page_id':page['id'],'properties':changes,'before':{k:page['properties'][k] for k in changes},'marker':marker,'history':history,'message_ids':d['messages']})
            if d.get('tracking'):
                tr=[p for p in tracking['results'] if (flat(p).get('邮箱') or '').lower()==d['email'].lower() or flat(p).get('达人账号名','').lstrip('@')==f['账号'].lstrip('@')]
                assert len(tr)<=1
                assert tr or d['tracking']=='upsert',(d['email'],'Expected existing tracking row')
                tracking_extra=dict(d.get('tracking_extra',{}))
                structured='合作选项' in ts
                if structured:
                    assert '名字（或昵称）' not in tracking_extra,'Use explicit contact_name or recipient_name; legacy names are ambiguous'
                    assert 'tracking_note' in d,'Structured tracking requires a concise tracking_note with next action and owner'
                    assert '佣金：' in d['tracking_note'],'Tracking notes must retain the evidence-based commission status'
                    for source,target in [('contact_name','邮件联系人姓名'),('recipient_name','收件人姓名'),('cooperation_option','合作选项'),('selected_product','选品')]:
                        if source in d:tracking_extra[target]=d[source]
                    if '合作选项' in tracking_extra:
                        assert tracking_extra['合作选项'] in {'Option 1','Option 2','无发布义务赠礼','未确认'}
                tv={'最新跟进日期':max(local_date(m['meta']) for m in ms),**tracking_extra,'建联备注':d['tracking_note'] if structured else note}
                if not tr:
                    tv.update({'达人账号名':f['账号'],'邮箱':d['email']})
                    if structured:
                        tv.setdefault('合作选项','未确认')
                        tv.setdefault('选品','未选品')
                    else:tv['名字（或昵称）']=d['tracking_name']
                if structured:
                    current_tracking=flat(tr[0]) if tr else {}
                    combined={**current_tracking,**tv}
                    option=combined.get('合作选项')
                    if isinstance(option,dict):option=option.get('name')
                    ready=option in {'Option 1','Option 2','无发布义务赠礼'} and all(combined.get(k) for k in ['收件人姓名','地址','电话','邮编'])
                    product=str(combined.get('选品') or '').strip()
                    selected=bool(product) and not any(word in product for word in ['未选品','待确认','待定','意向'])
                    if d['status']=='待寄样':
                        assert ready and selected, '待寄样 requires confirmed cooperation, recipient details and final product selection'
                    for field,default in [('是否已寄样','未寄样'),('是否已返图','未返图')]:
                        existing=current_tracking.get(field)
                        if isinstance(existing,dict):existing=existing.get('name')
                        # Preserve actual progress entered manually; never erase shipment evidence.
                        if existing and existing not in {'未寄样','未返图','待核实'}:continue
                        if not ready or not selected:tv[field]=None
                        elif field not in tv:
                            if not existing or existing=='待核实':tv[field]=default
                tp={k:typed(ts[k],v) for k,v in tv.items()}
                plan.append({'table':'tracking','handle':f['账号'],'page_id':tr[0]['id'] if tr else None,'properties':tp,'before':{k:tr[0]['properties'][k] for k in tp} if tr else {},'source_main_id':page['id'],'message_ids':d['messages']})
        save(self.o/'update_plan.json',plan)
        preview=[{'table':x['table'],'handle':x['handle'],'action':'update' if x['page_id'] else 'create','fields':list(x['properties'])} for x in plan]
        save(self.o/'update_preview.json',preview)
        print(json.dumps(preview,ensure_ascii=False),flush=True)
    def apply(self):
        plan=read(self.o/'update_plan.json');review=read(self.o/'review.json');assert review['plan_sha256']==digest(plan)
        needed=set()
        for x in plan:
            prop=x['properties'].get('联系状态')
            if prop:needed|=status_names(prop.get('multi_select',prop.get('select')))
        schema=self.r.notion('GET','data_sources/'+self.r.ds)
        kind=schema['properties']['联系状态']['type']
        options=schema['properties']['联系状态'][kind]['options'];missing=needed-{x['name'] for x in options}
        if missing:
            assert missing<={'待寄样'},'Unreviewed status option'
            updated_options=[{'id':x['id'],'name':x['name'],'color':x['color']} for x in options]+[{'name':name,'color':'purple'} for name in sorted(missing)]
            self.r.notion('PATCH','data_sources/'+self.r.ds,json={'properties':{'联系状态':{kind:{'options':updated_options}}}})
        log=read(self.o/'apply_log.json') if (self.o/'apply_log.json').exists() else {}
        for item in plan:
            key=item['table']+':'+item['handle'];state=log.get(key,{})
            if state.get('status')=='verified':continue
            pid=item['page_id'] or state.get('page_id')
            if not pid:
                # Re-query by exact email before creating, including after an uncertain POST.
                rows=self.r.notion('POST','data_sources/'+self.r.c['notion_tracking_data_source']+'/query',json={'filter':{'property':'邮箱','email':{'equals':item['properties']['邮箱']['email']}},'page_size':100})['results']
                assert len(rows)<=1
                if rows:pid=rows[0]['id']
            if pid:
                page=self.r.notion('GET','pages/'+pid)
                for k,v in item['properties'].items():
                    try:self.r.verify_properties(page,{k:v})
                    except AssertionError:
                        assert k in item['before'],'Unmatched existing tracking row'
                        old=item['before'][k];self.r.verify_properties(page,{k:{old['type']:old[old['type']]}})
                page=self.r.notion('PATCH','pages/'+pid,json={'properties':item['properties']})
            else:
                log[key]={'status':'creating','started_at':now()};save(self.o/'apply_log.json',log)
                page=self.r.notion('POST','pages',json={'parent':{'type':'data_source_id','data_source_id':self.r.c['notion_tracking_data_source']},'properties':item['properties'],'children':[para('达人主表：https://www.notion.so/'+item['source_main_id'].replace('-',''))]})
                pid=page['id'];log[key]={'status':'created','page_id':pid};save(self.o/'apply_log.json',log)
            self.r.verify_properties(page,item['properties'])
            if item.get('history'):
                blocks=self.r.blocks(pid)
                if not any(text(b)==item['marker'] for b in blocks):
                    self.r.notion('PATCH','blocks/'+pid+'/children',json={'children':item['history']})
            log[key]={'status':'verified','page_id':pid,'verified_at':now(),'url':page['url'],'message_ids':item['message_ids']};save(self.o/'apply_log.json',log)
            print('UPDATED',key,flush=True);time.sleep(.4)
    def verify(self):
        plan=read(self.o/'update_plan.json');log=read(self.o/'apply_log.json') if (self.o/'apply_log.json').exists() else {};assert len(log)==len(plan)
        conflicts=[];pages=[]
        main_pages={p['id']:p for p in self.r.all_pages()} if sum(x['table']=='main' for x in plan)>50 else {}
        for item in plan:
            pid=log[item['table']+':'+item['handle']]['page_id'];page=main_pages.get(pid) or self.r.notion('GET','pages/'+pid);pages.append(page)
            for key,value in item['properties'].items():
                try:self.r.verify_properties(page,{key:value})
                except AssertionError:
                    conflicts.append({'table':item['table'],'handle':item['handle'],'page_id':pid,'field':key,'expected':value,'actual':page['properties'][key],'last_edited_time':page['last_edited_time']})
        save(self.o/'verification_pages.json',pages);save(self.o/'verification_conflicts.json',conflicts)
        assert not conflicts,'Concurrent changes detected; inspect verification_conflicts.json before reconciling'
        save(self.o/'verified.json',{'verified_at':now(),'main_records':sum(x['table']=='main' for x in plan),'tracking_records':sum(x['table']=='tracking' for x in plan),'unique_messages':len({m for x in plan for m in x['message_ids']})})
        print(json.dumps(read(self.o/'verified.json'),ensure_ascii=False),flush=True)

    def export(self):
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        verified=read(self.o/'verified.json')
        plan=read(self.o/'update_plan.json');log=read(self.o/'apply_log.json')
        decisions=read(self.o/'decisions.json')
        messages={str(m['meta']['messageId']):m for m in read(self.o/'message_contents.json')}
        wb=Workbook();sheet=wb.active;sheet.title='后续跟进'
        sheet.append(['达人','姓名','往来邮箱','当前状态','暂停推进','合作方式','最新进展与报价','下一步','当天收到','当天跟进发出','当天最后来信（北京时间）','当天最后发出（北京时间）','Notion 主表','Notion 跟进表'])
        evidence=wb.create_sheet('邮件索引');evidence.append(['北京时间','达人','方向','发件人','收件人','主题','Zoho Message ID'])
        zone=timezone(timedelta(hours=8))
        stamp=lambda m:datetime.fromtimestamp(int(m['meta']['receivedTime'])/1000,zone).strftime('%Y-%m-%d %H:%M')
        owner={}
        for d in decisions:
            item=next(x for x in plan if x['table']=='main' and x['message_ids']==d['messages']);handle=item['handle']
            incoming=[messages[mid] for mid in d['messages'] if messages[mid]['direction']=='inbox']
            sent=[messages[mid] for mid in d['messages'] if messages[mid]['direction']=='sent']
            sheet.append([handle,d.get('tracking_name',''),d['email'],d['status'],'是' if d['paused'] else '否',d['cooperation'],d['summary'],d['next'],len(incoming),len(sent),max(map(stamp,incoming),default='当天无新来信'),max(map(stamp,sent),default='当天无跟进发出'),log['main:'+handle]['url'],log.get('tracking:'+handle,{}).get('url','')])
            owner.update({mid:handle for mid in d['messages']})
        for mid in sorted(owner,key=lambda mid:int(messages[mid]['meta']['receivedTime'])):
            m=messages[mid];meta=m['meta']
            evidence.append([stamp(m),owner[mid],{'inbox':'收到回复','sent':'发出跟进','system':'系统通知'}[m['direction']],str(meta.get('fromAddress','')),str(meta.get('toAddress','')),meta.get('subject',''),mid])
        scope=wb.create_sheet('统计口径')
        scope.append(['项目','说明']);scope.append(['范围',read(self.o/'collection_scope.json')['date']+' 00:00–24:00，北京时间'])
        scope.append(['内容','仅后续往来，不重复统计首次邀请。当天无新来信不代表历史从未回复。'])
        scope.append(['进展判断','报价不代表已接受；收货资料齐全不代表已寄样；无义务赠礼不代表承诺返图。'])
        scope.append(['核实结果',json.dumps(verified,ensure_ascii=False)])
        for ws in wb:
            ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
            for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='405D51')
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    c.alignment=Alignment(vertical='top',wrap_text=True)
                    if isinstance(c.value,str):c.data_type='s'
                ws.row_dimensions[row[0].row].height=150 if ws==sheet else 42
            for col in range(1,ws.max_column+1):ws.column_dimensions[get_column_letter(col)].width=24
        sheet.column_dimensions['G'].width=85;sheet.column_dimensions['H'].width=60
        sheet.column_dimensions['F'].width=40;evidence.column_dimensions['F'].width=65;scope.column_dimensions['B'].width=100
        path=self.o/('NUWARIA_'+read(self.o/'collection_scope.json')['date']+'_后续跟进.xlsx')
        wb.save(path)
        check=load_workbook(path,read_only=True)
        assert check['后续跟进'].max_row-1==verified['main_records']
        assert check['邮件索引'].max_row-1==verified['unique_messages']
        check.close();print('EXPORTED',path,flush=True)

if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--project',required=True);ap.add_argument('--folder',required=True);ap.add_argument('action',choices=['collect','attachments','prepare','apply','verify','export']);a=ap.parse_args();getattr(MailSync(a.project,a.folder),a.action)()
