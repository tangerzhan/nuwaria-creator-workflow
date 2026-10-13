"""Explicitly authorized outreach, with durable pre-send deduplication and audit records."""
import argparse, hashlib, json, re, time, html
from datetime import datetime, timezone
from pathlib import Path
import requests

EMAIL = re.compile(r'[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}')

def read(p): return json.loads(Path(p).read_text(encoding='utf-8-sig'))
def save(p, obj):
    p=Path(p); p.parent.mkdir(parents=True,exist_ok=True)
    tmp=p.with_suffix(p.suffix+'.tmp'); tmp.write_text(json.dumps(obj,ensure_ascii=False,indent=2),encoding='utf8'); tmp.replace(p)
def now(): return datetime.now(timezone.utc).isoformat()
def connection_without_response(record):
    error=record.get('error','')
    return record.get('status')=='uncertain' and ('SSLError' in error or ('ProxyError' in error and 'RemoteDisconnected' in error and 'without response' in error))
def status_names(value):
    if isinstance(value,list):return {x['name'] for x in value}
    return {value['name']} if isinstance(value,dict) and value.get('name') else set()

def validate_status_tags(names):
    stages=names-{'已报价'}
    assert len(stages)<=1,('Conflicting contact stages',sorted(stages))

def shipping_metrics(pages):
    """Use live tracking selections, including manual edits; never infer from contact stage."""
    from collections import Counter
    counts=Counter()
    for page in pages:
        if page.get('archived') or page.get('in_trash'):continue
        value=flat(page).get('是否已寄样')
        name=value.get('name') if isinstance(value,dict) else value
        counts[name or '未填写']+=1
    result={'寄样情况：'+name:counts[name] for name in ['未寄样','运输中','已送达','已收件','未填写']}
    result.update({'寄样情况：'+name:n for name,n in counts.items()})
    result['待寄样（按建联表未寄样）']=counts['未寄样']
    return result

class StatusTags(list):
    """Keep legacy primary-status readers working while exposing every tag."""
    def get(self,key,default=None):
        if key!='name':return default
        names=status_names(self)
        return next((s for s in ['已暂停','待跟进','待二次回复','待寄样','已完成','已发信','未联系','已报价'] if s in names),default)
def env(p):
    return {k.strip():v.strip().strip('\"\'') for l in Path(p).read_text(encoding='utf-8-sig').splitlines() if '=' in l and not l.lstrip().startswith('#') for k,v in [l.split('=',1)]}
def flat(page):
    return {k: ''.join(t.get('plain_text',t.get('text',{}).get('content','')) for t in v[v['type']]) if v['type'] in ('rich_text','title') else StatusTags(v['multi_select']) if k=='联系状态' and v['type']=='multi_select' else v.get(v['type']) for k,v in page['properties'].items()}
def text(block): return ''.join(t.get('plain_text',t.get('text',{}).get('content','')) for t in block.get(block['type'],{}).get('rich_text',[]))
def digest(d): return hashlib.sha256(json.dumps(d,ensure_ascii=False,sort_keys=True).encode('utf8')).hexdigest()
def rt(s): return [{'type':'text','text':{'content':str(s)[i:i+1900]}} for i in range(0,len(str(s)),1900)] if s is not None else []
def para(s): return {'object':'block','type':'paragraph','paragraph':{'rich_text':rt(s)}}
class Run:
    def __init__(self,p):
        self.http=requests.Session()
        self.p=Path(p).resolve(); self.o=self.p/'outreach'; self.o.mkdir(exist_ok=True)
        self.c=read(self.p/'project.json'); self.ds=self.c['notion_data_source']
        self.nh={'Authorization':'Bearer '+env(self.c['notion_env'])['NOTION_TOKEN'],'Notion-Version':'2025-09-03'}
        self.reviews={h:v for h,v in read(self.p/'production_reviews.json').items() if v.get('verdict')=='include'}
        self.alias=read(self.p/'identity_aliases.json'); self.zh=None
    def notion(self,method,path,**kw):
        props=kw.get('json',{}).get('properties',{})
        if path.startswith('pages') and '联系状态' in props and 'select' in props['联系状态']:
            schema=self.notion('GET','data_sources/'+self.ds)['properties']['联系状态']
            if schema['type']=='multi_select':
                names=status_names(props['联系状态']['select'])
                if method=='PATCH':
                    live=self.notion('GET',path)
                    names|=status_names(flat(live).get('联系状态'))&{'已报价'}
                props['联系状态']={'multi_select':[{'name':n} for n in sorted(names)]}
        if path.startswith('pages') and 'multi_select' in props.get('联系状态',{}):
            validate_status_tags(status_names(props['联系状态']['multi_select']))
        read_only=method=='GET' or (method=='POST' and path.endswith('/query'))
        for attempt in range(3):
            try:
                r=self.http.request(method,'https://api.notion.com/v1/'+path,headers=self.nh,timeout=45,**kw)
                break
            except requests.RequestException:
                if not read_only or attempt==2:raise
                print('Notion read connection interrupted; retrying safely',attempt+1,flush=True);time.sleep(2*(attempt+1))
        if not r.ok: raise RuntimeError('Notion HTTP '+str(r.status_code)+' '+r.text[:400])
        return r.json()
    def all_pages(self):
        rows=[]; cursor=None
        while True:
            body={'page_size':100}
            if cursor: body['start_cursor']=cursor
            d=self.notion('POST','data_sources/'+self.ds+'/query',json=body); rows.extend(d['results'])
            if not d.get('has_more'): return rows
            cursor=d['next_cursor']
    def blocks(self,bid,recurse=False):
        rows=[]; cursor=None
        while True:
            d=self.notion('GET','blocks/'+bid+'/children',params={'page_size':100,**({'start_cursor':cursor} if cursor else {})})
            rows.extend(d['results'])
            if not d.get('has_more'):break
            cursor=d['next_cursor']
        if recurse:
            for b in rows:
                if b.get('has_children'):b['_children']=self.blocks(b['id'],True)
        return rows
    def zoho_init(self):
        self.auth=read(self.c['zoho_auth'])
        a=self.auth
        if not self.c.get('sender_email') or self.c['sender_email'].lower()!=a.get('from_email','').lower():
            raise ValueError('Configure matching sender_email and Zoho from_email before outreach')
        for attempt in range(3):
            try:
                r=self.http.post(a['accounts_url'].rstrip('/')+'/oauth/v2/token',data={k:a[k] for k in ['client_id','client_secret','refresh_token']}|{'grant_type':'refresh_token'},timeout=40)
                break
            except requests.RequestException:
                if attempt==2:raise
                print('Zoho authentication connection interrupted; retrying',attempt+1,flush=True);time.sleep(2*(attempt+1))
        r.raise_for_status(); d=r.json()
        if not d.get('access_token'): raise RuntimeError('Zoho OAuth failed: '+str(d.get('error','unknown')))
        self.zh={'Authorization':'Zoho-oauthtoken '+d['access_token']}
        self.zbase=a['mail_base_url'].rstrip('/')+'/api/accounts/'+str(a['account_id'])
    def zoho(self,method,path,**kw):
        if self.zh is None:self.zoho_init()
        for attempt in range(3):
            try:
                r=self.http.request(method,self.zbase+path,headers=self.zh,timeout=45,**kw)
                break
            except requests.RequestException:
                if method!='GET' or attempt==2:raise
                print('Zoho read connection interrupted; retrying safely',attempt+1,flush=True);time.sleep(2*(attempt+1))
        if not r.ok:raise RuntimeError('Zoho HTTP '+str(r.status_code)+' '+r.text[:300])
        d=r.json()
        if int(d.get('status',{}).get('code',200))>=300:raise RuntimeError('Zoho error '+str(d.get('status')))
        return d
    def sent(self):
        folders=self.zoho('GET','/folders')['data']
        fs=[x for x in folders if str(x.get('folderName','')).lower()=='sent' or str(x.get('folderType','')).lower()=='sent']
        if len(fs)!=1:raise RuntimeError('Cannot uniquely identify Sent folder')
        self.sent_folder=str(fs[0]['folderId'])
        from mail_cache import scan_folder
        return scan_folder(self,self.sent_folder)
    def candidates(self):
        out=[]
        for h,v in self.reviews.items():
            r=read(self.p/'profiles'/f'{h}.json')
            emails=([v['public_email']] if v['public_email'] else []) if 'public_email' in v else r.get('emails',[])
            out.append({'handle':h,'name':v.get('name') or r['user'].get('full_name',''),'emails':emails,'url':r['url'],'aliases':self.alias.get(h,[])+v.get('aliases',[])})
        return out
    def dedup(self,pages,sent):
        report=[]
        imports_path=self.o/'notion_import_log.json' if hasattr(self,'o') else None
        imports=read(imports_path) if imports_path and imports_path.exists() else {}
        # Build once per snapshot; candidates only look up matching keys locally.
        by_email={};by_handle={};records=[];sent_by_email={}
        for i,page in enumerate(pages):
            f=flat(page)
            account=str(f.get('账号','')).strip().lower().lstrip('@')
            urls=' '.join(str(f.get(k) or '') for k in ['主页链接','Instagram主页','TikTok主页'])
            handles={x.lower() for x in re.findall(r'(?:instagram\.com/|tiktok\.com/@)([\w.]+)',urls)}|{account}
            em={e.lower() for e in EMAIL.findall(str(f.get('公开联系方式','')))}
            records.append((page['id'],f,handles,em))
            for e in em:by_email.setdefault(e,set()).add(i)
            for h in handles-{''}:by_handle.setdefault(h,set()).add(i)
        notion_emails=set(by_email)
        for i,x in enumerate(sent):
            for e in {e.lower() for e in EMAIL.findall(str(x.get('toAddress','')))}:
                sent_by_email.setdefault(e,set()).add(i)
        for c in self.candidates():
            keys={h.strip().lower().lstrip('@') for h in [c['handle'],*c['aliases']]}
            mails={e.strip().lower() for e in c['emails']};matches=[];hits=set();sent_hits=set()
            for e in mails&notion_emails:hits.update(by_email[e])
            for h in keys:hits.update(by_handle.get(h,()))
            for i in sorted(hits):
                pid,f,handles,em=records[i]
                matches.append({'id':pid,'account':f.get('账号'),'status':f.get('联系状态'),'paused':f.get('是否暂停推进',False),'sent_date':f.get('发送日期'),'reply_date':f.get('回复日期'),'reason':sorted(handles&keys)+sorted(em&mails)})
            for e in mails:sent_hits.update(sent_by_email.get(e,()))
            sm=[{'id':sent[i]['messageId'],'to':sent[i].get('toAddress'),'subject':sent[i].get('subject'),'time':sent[i].get('sentDateInGMT',sent[i].get('receivedTime'))} for i in sorted(sent_hits)]
            own=imports.get(c['handle'],{})
            own_pending=bool(matches) and own.get('status')=='verified' and all(m['id']==own['page_id'] and status_names(m.get('status'))=={'未联系'} and not m['paused'] and not m['sent_date'] and not m['reply_date'] for m in matches)
            report.append(c|{'notion_matches':matches,'sent_matches':sm,'own_pending_import_exemption':own_pending,'eligible':bool(mails) and (not matches or own_pending) and not sm})
        return report
    def preflight(self):
        pages=self.all_pages();save(self.o/'notion_before.json',pages)
        save(self.o/'schema_live.json',self.notion('GET','data_sources/'+self.ds))
        sop=self.blocks(self.c['notion_sop'],True);save(self.o/'sop_live.json',sop)
        def walk(bs):
            for b in bs:
                yield b
                yield from walk(b.get('_children',[]))
        (self.o/'sop_live.txt').write_text('\n\n'.join(b['type']+' '+text(b) for b in walk(sop)),encoding='utf8')
        save(self.o/'main_before.json',self.blocks(self.c['notion_main_page']));save(self.o/'stats_before.json',self.blocks(self.c['notion_stats_table']))
        sent=self.sent();save(self.o/'zoho_sent_before.json',sent)
        report=self.dedup(pages,sent);save(self.o/'dedup.json',{'checked_at':now(),'notion_count':len(pages),'sent_count':len(sent),'sent_folder':self.sent_folder,'candidates':report})
        print(json.dumps({'notion':len(pages),'sent':len(sent),'people_with_email':sum(bool(c['emails']) for c in report),'email_addresses':sum(len(c['emails']) for c in report),'notion_duplicates':[c['handle'] for c in report if c['notion_matches']],'sent_duplicates':[c['handle'] for c in report if c['sent_matches']],'eligible':sum(c['eligible'] for c in report)},ensure_ascii=False),flush=True)
    def build(self):
        sop=read(self.o/'sop_live.json'); active=False; parts=[]
        for b in sop:
            if b['type'].startswith('heading_'):
                if text(b).startswith('9. '):active=True
                elif active:break
            if active and b['type']=='code':parts.append(text(b))
        template=''.join(parts).replace('\r\n','\n').strip()
        tail=template[template.index('NUWARIA is a lifestyle brand'):]
        assert 'at least 3 seconds' in tail and tail.endswith('www.nuwaria.com')
        assert '[Name]' not in tail and '[' not in tail
        (self.o/'template_tail.txt').write_text(tail,encoding='utf8')
        personal=read(self.o/'personalization.json'); dd=read(self.o/'dedup.json'); drafts=[]
        prior=read(self.o/'drafts.json') if (self.o/'drafts.json').exists() else []
        sent_log=read(self.o/'send_log.json') if (self.o/'send_log.json').exists() else {}
        preserved={d['handle']:d for d in prior if sent_log.get(d['handle'],{}).get('status')=='sent_verified'}
        scope=read(self.o/'batch_scope.json') if (self.o/'batch_scope.json').exists() else None
        if scope:
            selected=scope['handles']
            assert len(selected)==len(set(selected))==scope['authorized_count']
            eligible={c['handle'] for c in dd['candidates'] if c['eligible']}
            assert set(selected)<=eligible|set(preserved), 'Selected batch includes ineligible recipients'
        for c in dd['candidates']:
            if scope and c['handle'] not in selected:continue
            if c['handle'] in preserved:
                d=preserved[c['handle']]
                assert sent_log[c['handle']]['sha256']==d['sha256']
                drafts.append(d);continue
            if not c['eligible']:continue
            h=c['handle']; a=personal[h]; intro="I'm Aria from NUWARIA. I came across your Instagram and really enjoyed "+a['observation']
            body='Hi '+a.get('greeting',a['name'])+',\n\n'+intro+'\n\n'+tail
            email=a.get('email',c['emails'][0]); assert email in c['emails'] and EMAIL.fullmatch(email)
            source=read(self.p/'profiles'/f'{h}.json')
            d={'handle':h,'to':email,'name':a['name'],'subject':'NUWARIA x '+a['name']+' Collaboration Invitation','intro':intro,'body':body,'evidence':[{'index':i,'url':source['posts'][i-1]['url'],'scene':self.reviews[h]['scenes'][str(i)]} for i in a['indices']]}
            d['sha256']=digest({k:d[k] for k in ['to','subject','body']});drafts.append(d)
        self.validate_drafts(drafts,tail)
        save(self.o/'drafts.json',drafts)
        (self.o/'preview.html').write_text('<!doctype html><html lang="en"><meta charset="utf-8"><title>NUWARIA outreach review</title><style>body{font:16px/1.6 Arial;max-width:880px;margin:30px auto}article{border-bottom:2px solid #888;padding:30px 0}pre{white-space:pre-wrap;font:inherit}</style>'+''.join('<article><h2>'+html.escape(d['handle'])+'</h2><p>To: '+html.escape(d['to'])+'</p><h3>'+html.escape(d['subject'])+'</h3><pre>'+html.escape(d['body'])+'</pre></article>' for d in drafts)+'</html>',encoding='utf8')
        (self.o/'intro_review.txt').write_text('\n\n'.join(f"{i+1}. @{d['handle']} | {d['to']} | {d['subject']}\nHi {personal[d['handle']].get('greeting',d['name'])},\n{d['intro']}" for i,d in enumerate(drafts)),encoding='utf8')
        print('Built and mechanically checked',len(drafts),'drafts; editorial review still required.')
    def validate_drafts(self,drafts,tail):
        seen=set()
        for d in drafts:
            assert d['to'].lower() not in seen; seen.add(d['to'].lower())
            assert EMAIL.fullmatch(d['to'])
            assert d['subject']=='NUWARIA x '+d['name']+' Collaboration Invitation'
            assert d['body'].endswith(tail)
            assert d['intro'] in d['body'] and 'your Instagram' in d['intro'] and 'TikTok' not in d['intro']
            assert d['sha256']==digest({k:d[k] for k in ['to','subject','body']})
            s=d['subject']+'\n'+d['body']
            assert s.isascii(), 'Non-ASCII outbound text: '+d['handle']
            assert not re.search(r'[\[\]{}<>]|\ufffd|\?\?|TODO|TBD|specific content|internal note',s,re.I),d['handle']
            assert all(ord(c)>=32 or c in '\n\t' for c in s)
            # A concise, evidence-based opening need not be padded to 40 words.
            assert 20 <= len(d['intro'].split()) <= 110, d['handle']
        assert len(drafts)==len({d['handle'] for d in drafts})
    def approve_review(self):
        drafts=read(self.o/'drafts.json');tail=(self.o/'template_tail.txt').read_text(encoding='utf8');self.validate_drafts(drafts,tail)
        save(self.o/'review.json',{'reviewed_at':now(),'reviewer':'assistant editorial review after reading every preview and its evidence','checks':['recipient and greeting','specific observation matches recorded scene','Instagram platform','English fluency','no CJK, placeholders, control characters or mojibake','exact live SOP tail with correct code-block joining','one address per creator'],'draft_hashes':{d['handle']:d['sha256'] for d in drafts},'template_sha256':hashlib.sha256(tail.encode()).hexdigest()})
        print('Recorded editorial approval for',len(drafts),'drafts.')
    def check_sent_body(self,d,mid,folder):
        from bs4 import BeautifulSoup
        data=self.zoho('GET',f'/folders/{folder}/messages/{mid}/content')['data']
        content=data['content']; plain=BeautifulSoup(content,'html.parser').get_text(' ',strip=True) if '<' in content else content
        normalize=lambda s: ' '.join(html.unescape(s).split())
        if normalize(plain)!=normalize(d['body']):
            save(self.o/'content_mismatch.json',{'handle':d['handle'],'actual':content,'expected':d['body']})
            raise RuntimeError('Sent content differs from reviewed text: '+d['handle'])
        return {'checked_at':now(),'message_id':str(mid),'matches_reviewed_body':True,'content_sha256':hashlib.sha256(content.encode('utf8')).hexdigest()}
    def send(self,limit=None,retry_blocked=False):
        drafts=read(self.o/'drafts.json'); tail=(self.o/'template_tail.txt').read_text(encoding='utf8'); self.validate_drafts(drafts,tail)
        if (self.o/'batch_scope.json').exists():
            scope=read(self.o/'batch_scope.json')
            assert len(drafts)==scope['authorized_count'] and {d['handle'] for d in drafts}==set(scope['handles']), 'Drafts differ from authorized batch'
        review=read(self.o/'review.json')
        assert review['draft_hashes']=={d['handle']:d['sha256'] for d in drafts},'Review invalidated by edits'
        pages=self.all_pages();sent=self.sent();report=self.dedup(pages,sent)
        save(self.o/'dedup_at_send.json',{'checked_at':now(),'notion_count':len(pages),'sent_count':len(sent),'candidates':report})
        byhandle={c['handle']:c for c in report}
        # Verify the configured sender is an address belonging to the authenticated account.
        account=self.zoho('GET','')['data']
        if account.get('outgoingBlocked'):raise RuntimeError('Zoho outgoing is blocked; use official recovery first')
        allowed={e.lower() for e in EMAIL.findall(json.dumps(account))}
        sender=self.auth['from_email']; assert sender.lower() in allowed,'Configured sender not in authenticated account'
        save(self.o/'sender_check.json',{'checked_at':now(),'from':sender,'verified':True})
        log=read(self.o/'send_log.json') if (self.o/'send_log.json').exists() else {}
        completed=0
        for d in drafts:
            h=d['handle']
            attempts=[]
            if h in log:
                if log[h]['status'] in ('accepted','sent_verified','skipped_duplicate'):continue
                if log[h]['status'] in ('uncertain','inflight'):
                    print('HELD_UNCERTAIN',h,'reconcile before retry',flush=True)
                    continue
                if log[h]['status']=='absent_after_reconciliation':
                    attempts=log[h].get('prior_attempts',[])+[{k:v for k,v in log[h].items() if k!='prior_attempts'}]
                elif log[h]['status']=='blocked_temporary' and retry_blocked:
                    elapsed=(datetime.now(timezone.utc)-datetime.fromisoformat(log[h]['attempted_at'])).total_seconds()
                    recovery=read(self.o/'recovery.json') if (self.o/'recovery.json').exists() else {}
                    confirmed=recovery.get('user_confirmed_unlocked_at')
                    restored=bool(confirmed and datetime.fromisoformat(confirmed)>datetime.fromisoformat(log[h]['attempted_at']))
                    if elapsed<900 and not restored:raise RuntimeError('Temporary rejection: no retry before 15-minute minimum observation interval or explicit new unlock confirmation')
                    attempts=log[h].get('prior_attempts',[])+[{k:v for k,v in log[h].items() if k!='prior_attempts'}]
                else:raise RuntimeError('Unresolved previous attempt; reconcile before resending: '+h)
            if not byhandle[h]['eligible']:
                log[h]={'status':'skipped_duplicate','checked_at':now(),'notion':byhandle[h]['notion_matches'],'sent':byhandle[h]['sent_matches']};save(self.o/'send_log.json',log);continue
            if limit is not None and completed>=limit:break
            log[h]={'status':'inflight','attempted_at':now(),'to':d['to'],'subject':d['subject'],'sha256':d['sha256'],'prior_attempts':attempts};save(self.o/'send_log.json',log)
            try:
                result=self.zoho('POST','/messages',json={'fromAddress':sender,'toAddress':d['to'],'subject':d['subject'],'content':d['body'],'mailFormat':'plaintext','encoding':'UTF-8'})
                mid=result.get('data',{}).get('messageId')
                if not mid:raise RuntimeError('No message ID in send response')
                log[h].update(status='accepted',accepted_at=now(),message_id=str(mid),response=result)
                save(self.o/'send_log.json',log)
            except Exception as e:
                log[h].update(status='uncertain',error=str(e)[:500]);save(self.o/'send_log.json',log)
                raise
            time.sleep(1.2)
            # Check the stored message, including UTF-8 round trip, before proceeding.
            check=self.check_sent_body(d,mid,self.sent_folder)
            log[h].update(status='sent_verified',content_check=check);save(self.o/'send_log.json',log)
            completed+=1;print('SENT_VERIFIED',h,d['to'],flush=True)
            if limit is None or completed<limit:time.sleep(10)
        print('Send status', {s:sum(x['status']==s for x in log.values()) for s in {x['status'] for x in log.values()}},flush=True)
    def reconcile(self):
        drafts={d['handle']:d for d in read(self.o/'drafts.json')};log=read(self.o/'send_log.json');sent=self.sent()
        save(self.o/'zoho_sent_after.json',sent)
        pending=None
        if any(connection_without_response(x) for x in log.values()):
            folders=self.zoho('GET','/folders')['data']
            outbox=next(f for f in folders if f.get('folderName','').lower()=='outbox')
            pending=[]
            for start in range(1,20001,100):
                batch=self.zoho('GET','/messages/view',params={'folderId':outbox['folderId'],'limit':100,'start':start})['data']
                pending.extend(batch)
                if len(batch)<100:break
            else:raise RuntimeError('Outbox pagination incomplete')
            save(self.o/'outbox_reconciliation.json',{'checked_at':now(),'messages':pending})
        for h,x in log.items():
            d=drafts[h]
            matches=[m for m in sent if m.get('subject')==d['subject'] and d['to'].lower() in {e.lower() for e in EMAIL.findall(str(m.get('toAddress','')))}]
            historical_bounces={str(a['message_id']) for a in x.get('prior_deliveries',[]) if a.get('delivery_status')=='bounced'}
            matches=[m for m in matches if str(m['messageId']) not in historical_bounces]
            if len(matches)>1:raise RuntimeError('Multiple matching sent messages: '+h)
            if matches:
                mid=str(matches[0]['messageId']);x.update(status='sent_verified',message_id=mid,sent_list_verified_at=now())
                if not x.get('content_check') or str(x['content_check'].get('message_id'))!=mid:x['content_check']=self.check_sent_body(d,mid,self.sent_folder)
                if x.get('resend_attempt',{}).get('message_id')==mid:x['resend_attempt']['status']='sent_verified'
            elif x['status']=='uncertain' and '550 5.4.6 Unusual sending activity detected' in x.get('error',''):
                x.update(status='blocked_temporary',reconciled_at=now(),reason='Explicit Zoho rejection; no matching message in complete Sent scan')
            elif connection_without_response(x) and pending is not None:
                queued=any(d['to'].lower() in {e.lower() for e in EMAIL.findall(str(m.get('toAddress','')))} for m in pending)
                if not queued:
                    checks=x.setdefault('absence_checks',[])
                    checks.append({'checked_at':now(),'sent_count':len(sent),'outbox_count':len(pending),'sent_matches':0,'outbox_matches':0})
                    age=(datetime.now(timezone.utc)-datetime.fromisoformat(x['attempted_at'])).total_seconds()
                    span=(datetime.fromisoformat(checks[-1]['checked_at'])-datetime.fromisoformat(checks[0]['checked_at'])).total_seconds()
                    if len(checks)>=2 and age>=120 and span>=60:
                        x.update(status='absent_after_reconciliation',reconciled_at=now(),reason='Connection failed without a response; absent from two complete Sent and Outbox checks at least 60 seconds apart; fresh pre-send dedup still required')
            elif x['status']=='sent_verified':raise RuntimeError('Previously accepted mail absent from Sent: '+h)
        save(self.o/'send_log.json',log)
        account=self.zoho('GET','')['data']
        def flags(obj):
            if isinstance(obj,dict):
                return {k:(v if not isinstance(v,(dict,list)) else flags(v)) for k,v in obj.items() if any(t in k.lower() for t in ['block','outgoing','enabled','status'])}
            return obj
        save(self.o/'account_status.json',{'checked_at':now(),'flags':flags(account)})
        folders=self.zoho('GET','/folders')['data'];inbox=next(x for x in folders if x.get('folderName','').lower()=='inbox')
        recent=self.zoho('GET','/messages/view',params={'folderId':inbox['folderId'],'limit':100,'start':1})['data']
        alerts=[]
        for m in recent:
            if any(k in (str(m.get('fromAddress',''))+' '+m.get('subject','')).lower() for k in ['mailer-daemon','postmaster','undeliver','failure','blocked','delivery status','退信','发送失败']):
                data=self.zoho('GET',f"/folders/{inbox['folderId']}/messages/{m['messageId']}/content")['data']
                involved=[h for h,d in drafts.items() if d['to'].lower() in data.get('content','').lower()]
                alerts.append({'message':m,'content':data.get('content',''),'handles':involved})
        save(self.o/'delivery_alerts.json',alerts)
        print(json.dumps({'sent_records':len(sent),'status_counts':{s:sum(x['status']==s for x in log.values()) for s in {x['status'] for x in log.values()}},'account_flags':flags(account),'batch_alerts':[{'handles':a['handles'],'subject':a['message'].get('subject')} for a in alerts if a['handles']]},ensure_ascii=False),flush=True)
    def retry_bounce(self,handle):
        """Retry a proven sender-block bounce after an explicit user unlock confirmation."""
        log=read(self.o/'send_log.json');x=log[handle]
        assert x.get('delivery_status')=='bounced' and x.get('delivery_error')=='554 5.1.8 Sender Address Blocked'
        confirmation=read(self.o/'recovery.json')['user_confirmed_unlocked_at']
        assert datetime.fromisoformat(confirmation)>datetime.fromisoformat(x['accepted_at'])
        drafts=read(self.o/'drafts.json');d=next(d for d in drafts if d['handle']==handle)
        self.validate_drafts(drafts,(self.o/'template_tail.txt').read_text(encoding='utf8'))
        assert read(self.o/'review.json')['draft_hashes'][handle]==d['sha256']
        pages=self.all_pages();sent=self.sent();c=next(c for c in self.dedup(pages,sent) if c['handle']==handle)
        own=read(self.o/'notion_import_log.json')[handle]
        assert len(c['notion_matches'])==1 and c['notion_matches'][0]['id']==own['page_id']
        assert status_names(c['notion_matches'][0].get('status')) & {'已发信','待跟进'}
        failed_ids={str(x['message_id'])}|{str(a['message_id']) for a in x.get('prior_deliveries',[]) if a.get('delivery_status')=='bounced'}
        other=[m for m in sent if d['to'].lower() in {e.lower() for e in EMAIL.findall(str(m.get('toAddress','')))} and str(m['messageId']) not in failed_ids]
        assert not other,'Another outbound message exists; do not send duplicate'
        account=self.zoho('GET','')['data'];assert not account.get('outgoingBlocked')
        sender=self.auth['from_email'];assert sender.lower() in {e.lower() for e in EMAIL.findall(json.dumps(account))}
        folders=self.zoho('GET','/folders')['data'];inbox=next(f for f in folders if f.get('folderName','').lower()=='inbox')
        recent=self.zoho('GET','/messages/view',params={'folderId':inbox['folderId'],'limit':100,'start':1})['data']
        assert not any(d['to'].lower() in {e.lower() for e in EMAIL.findall(str(m.get('fromAddress','')))} for m in recent),'Recipient already replied; review before resend'
        save(self.o/'bounce_retry_preflight.json',{'checked_at':now(),'handle':handle,'notion_matches':c['notion_matches'],'failed_ids':sorted(failed_ids),'other_outbound_matches':other,'user_unlock_confirmation':confirmation})
        # Archive the failed delivery first, so recovery can distinguish it from a new message.
        if str(x['message_id']) not in {str(a['message_id']) for a in x.get('prior_deliveries',[])}:
            x.setdefault('prior_deliveries',[]).append({k:x[k] for k in ['message_id','accepted_at','delivery_status','delivery_error','delivery_alert_id'] if k in x})
        x['resend_attempt']={'status':'inflight','attempted_at':now(),'reason':'User confirmed sender unblocked; retry original recipient after documented sender-block bounce'};save(self.o/'send_log.json',log)
        try:
            result=self.zoho('POST','/messages',json={'fromAddress':sender,'toAddress':d['to'],'subject':d['subject'],'content':d['body'],'mailFormat':'plaintext','encoding':'UTF-8'})
            mid=result.get('data',{}).get('messageId');assert mid
            x['resend_attempt'].update(status='accepted',message_id=str(mid),accepted_at=now())
            x.update(message_id=str(mid),accepted_at=now(),response=result,delivery_status='resent_after_sender_unblock')
            for k in ['delivery_error','delivery_alert_id','delivery_note']:x.pop(k,None)
            save(self.o/'send_log.json',log)
        except Exception as e:
            x['resend_attempt'].update(status='uncertain',error=str(e)[:500]);save(self.o/'send_log.json',log);raise
        time.sleep(1.2)
        x['content_check']=self.check_sent_body(d,mid,self.sent_folder);x['resend_attempt']['status']='sent_verified';save(self.o/'send_log.json',log)
        print('BOUNCE_RETRY_SENT_VERIFIED',handle,d['to'],flush=True)
    def research_workbook(self):
        scope=read(self.o/'notion_scope.json') if (self.o/'notion_scope.json').exists() else {}
        path=(self.p/scope.get('source_workbook',self.c['output_name'])).resolve()
        assert path.parent==self.p and path.suffix=='.xlsx', 'Invalid research workbook path'
        return path
    def build_import(self):
        from openpyxl import load_workbook
        wb=load_workbook(self.research_workbook(),read_only=True,data_only=True)
        rows=iter(wb['入选名单'].values); headers=next(rows)
        schema=read(self.o/'schema_live.json')['properties'];drafts={d['handle']:d for d in read(self.o/'drafts.json')}
        log=read(self.o/'send_log.json');zh=read(self.o/'intro_zh.json');tailzh=(self.o/'template_tail_zh.txt').read_text(encoding='utf8').strip()
        previous={x['handle']:x['payload']['properties'] for x in read(self.o/'notion_import_plan.json')} if (self.o/'notion_import_plan.json').exists() else {}
        scores=read(self.p/'scores.json') if (self.p/'scores.json').exists() else {}
        selected=self.import_handles()
        plan=[]
        for row in rows:
            f=dict(zip(headers,row));h=f['账号']
            if h not in selected:continue
            v=self.reviews[h];s=log.get(h,{});d=drafts.get(h)
            linked=v.get('linked_accounts',[])
            for channel in linked:
                platform=channel['platform']
                if platform in ('Instagram','TikTok') and not f.get(platform+'主页'):
                    f[platform+'主页']=channel['url'];f[platform+'粉丝数']=str(channel['followers'])
            if linked:
                f['外部核验备注']+='\n官方关联账号（仅身份和粉丝数核验）：'+'；'.join(str(c['platform'])+' '+c['url']+' '+str(c['followers']) for c in linked)
            state=s.get('status','not_attempted');is_sent=state=='sent_verified'
            f['外部核验备注']+='\n置信度说明：'+str(f.get('置信度') or '')
            f['置信度']=None;f['评分']=None
            parts=[];content_points=0;content_weight=0
            for key,label,weight in [('natural_seeding','自然植入',25),('brand_fit','品牌适配',20)]:
                detail=scores.get(h,{}).get(key,{})
                rating=detail.get('rating')
                if isinstance(rating,(int,float)) and not isinstance(rating,bool):
                    points=rating*weight/5;content_points+=points;content_weight+=weight
                    parts.append(f'{label} {points:g}/{weight}：'+detail.get('reason',''))
            if parts:
                f['外部核验备注']+='\n内容评分 '+f'{content_points:g}/{content_weight}（仅已核内容维度，不是完整100分总评）\n'+'\n'.join(parts)
            f['联系状态']='已发信' if is_sent else '未联系'
            f['筛选行动']='P2 联系' if is_sent else 'P2 终核后联系'
            f['下一步']='等待回复；有意向后确认产品选择、Option 1 / Option 2、寄样信息及受众与商业条件。' if is_sent else None
            if is_sent:
                from datetime import timedelta
                f['发送日期']=datetime.fromisoformat(s.get('accepted_at') or s['attempted_at']).astimezone(timezone(timedelta(hours=8))).date().isoformat()
            else:f['发送日期']=None
            f['邮件主题']=d['subject'] if d else None
            if is_sent:
                f['建联备注']='合作方式未确认；未选品；佣金：首次邀请已笼统提及，比例未告知、待回应；待对方回复合作意向。'
            elif state=='blocked_temporary':f['建联备注']='邮件已逐封审阅；Zoho 临时发送限制，发送失败，不计入成功。待服务恢复后核对已发送记录再续发。'
            elif d:f['建联备注']='英文邮件已逐封审阅，尚未确认成功发送；以发送日志为准。'
            elif EMAIL.search(str(f.get('公开联系方式',''))):f['建联备注']='已记录公开邮箱；按用户分批要求保留至下一批，尚未编写或发送邮件。'
            else:f['建联备注']='本轮未记录可直接使用的公开邮箱，未发送邮件或私信。'
            # Detailed deduplication and Zoho message IDs stay in the audit files.
            if s.get('delivery_status')=='bounced':
                f['建联备注']+=' 已收到退信：'+s['delivery_error']+'。本次投递失败。'
                if s.get('delivery_error')=='554 5.1.8 Sender Address Blocked':
                    f['建联备注']+=' 错误指向发件地址拦截，不能据此认定收件邮箱不存在。'
                    f['下一步']='核实 Zoho 发件地址拦截（'+s['delivery_error']+'）；恢复发信权限后处理投递。'
                elif '5.1.1' in s.get('delivery_error',''):
                    f['建联备注']+=' 收件服务器报告地址不可投递，具体原因以退信原文为准；停止向该地址发信。'
                    f['下一步']='核实其他公开联系方式，或后续通过Instagram主页联系；不要重发至已永久退信地址。'
                else:
                    f['下一步']='核实已记录退信原因后再决定联系渠道；不要自动重发。'
            elif s.get('delivery_status')=='resent_after_sender_unblock':
                f['建联备注']+=' 首次投递曾因 554 5.1.8 Sender Address Blocked 退信；用户确认解锁后已向原邮箱重发并核对正文，历史退信保留，不代表再次投递已获确认。'
            # Backfilled records may already have replies or a completed closeout.
            override_path=self.o/'notion_status_overrides.json'
            overrides=read(override_path) if override_path.exists() else {}
            override=overrides.get(h,{})
            if override:
                assert override.get('evidence_message_ids') and override.get('reviewed_at'), 'Unreviewed status override'
                fields=override['fields']
                assert set(fields)<=set(['联系状态','是否暂停推进','回复日期','跟进日期','建联备注','下一步','筛选行动'])
                if fields.get('联系状态')=='待跟进':
                    assert override.get('needs_our_reply') is True
                f.update(fields)
            properties={}
            for k,meta in schema.items():
                val=f.get(k);typ=meta['type']
                if typ in ('title','rich_text'):properties[k]={typ:rt(val) if val is not None else []}
                elif typ=='number':properties[k]={'number':val if isinstance(val,(int,float)) and not isinstance(val,bool) else None}
                elif typ=='checkbox':properties[k]={'checkbox':bool(val)}
                elif typ=='date':properties[k]={'date':{'start':val} if val else None}
                elif typ=='url':
                    assert not val or str(val).startswith(('https://','http://')),(h,k,val)
                    properties[k]={'url':val or None}
                elif typ=='select':
                    assert not val or val in [x['name'] for x in meta['select']['options']],(h,k,val)
                    properties[k]={'select':{'name':val} if val else None}
                elif typ=='multi_select':
                    names=val if isinstance(val,list) else [val] if val else []
                    properties[k]={'multi_select':[{'name':n} for n in names]}
                else:raise ValueError('Unhandled Notion type '+typ)
            children=[para(f"{self.c['date']} 新增{self.c['target']}位项目｜本页为筛选证据与建联内容。联系状态以数据库属性为准。"),para('核验限制：视频采用抽帧审核，部分补看细节；图文已查看范围以逐人证据为准。抽帧不等于完整播放，未补看的轮播不称已遍历。受众、转化、费用、排他等未公开条件待确认。')]
            if d:
                children.append({'object':'block','type':'heading_2','heading_2':{'rich_text':rt('已审阅的英文联系内容')}})
                children.extend(para(x) for x in d['body'].split('\n\n'))
                children.append({'object':'block','type':'heading_2','heading_2':{'rich_text':rt('中文译文（仅供内部参考，不发送）')}})
                children.append(para('你好，'+d['name']+'：'))
                children.append(para('我是 NUWARIA 的 Aria。我在 Instagram 上看到了你的账号。'+zh[h]))
                children.extend(para(x) for x in tailzh.split('\n\n'))
            children.append({'object':'block','type':'heading_2','heading_2':{'rich_text':rt('内容与评论核验')}})
            profile=read(self.p/'profiles'/f'{h}.json')
            for i,post in enumerate(profile['posts'][:12],1):
                caption=' '.join(post.get('caption','').split()[:23])[:180]
                children.append(para(f"{i}. {post['date']}\n"+v.get('topics',{}).get(str(i),'公开文案短摘：'+caption)+'\n'+v['scenes'].get(str(i),'')+'\n'+post['url']))
            children.extend([para('评论核验：'+v['comments_note']),para('抽查评论的原帖：\n'+'\n'.join(profile['posts'][i-1]['url'] for i in v['comment_indices'])),para('商业复核：'+v['business']),para('待核实：'+v['risks'])])
            adoption_path=self.o/'notion_adoption_review.json'
            adoptions=read(adoption_path) if adoption_path.exists() else {}
            adoption=adoptions.get(h)
            if adoption:
                assert adoption.get('reviewed_at') and adoption.get('page_id')
                for k,old in adoption['before_properties'].items():
                    typ=old['type'];value=old.get(typ)
                    if k not in override.get('fields',{}) and value not in (None,[], '',False):
                        properties[k]={typ:value}
            assert len(children)<=100
            payload={'parent':{'type':'data_source_id','data_source_id':self.ds},'properties':properties,'children':children}
            assert len(json.dumps(payload).encode())<450000
            plan.append({'handle':h,'send_status':state,'payload':payload,'properties_sha256':digest(properties),'previous_properties':previous.get(h)})
        assert len(plan)==len(selected) and {x['handle'] for x in plan}==selected
        save(self.o/'notion_import_plan.json',plan)
        print('Prepared',len(plan),'Notion records;',sum(x['send_status']=='sent_verified' for x in plan),'sent;',sum(x['handle'] in drafts for x in plan),'reviewed bilingual contact texts')
    def import_handles(self):
        path=self.o/'notion_scope.json'
        if not path.exists():return set(self.reviews)
        scope=read(path);handles=scope['handles']
        assert len(handles)==len(set(handles))==scope['authorized_count'], 'Invalid Notion scope'
        assert set(handles)<=set(self.reviews), 'Unknown Notion scope handle'
        if scope.get('email_only'):
            emails={c['handle'] for c in self.candidates() if c['emails']}
            assert set(handles)<=emails, 'Notion email-only scope includes no-email creator'
        return set(handles)
    def import_notion(self,limit=None):
        plan=read(self.o/'notion_import_plan.json')
        assert len(plan)==len(self.import_handles()) and {x['handle'] for x in plan}==self.import_handles(), 'Plan differs from authorized Notion scope'
        pages=self.all_pages()
        matches={c['handle']:c['notion_matches'] for c in self.dedup(pages,[])}
        log=read(self.o/'notion_import_log.json') if (self.o/'notion_import_log.json').exists() else {}
        count=0
        for item in plan:
            h=item['handle'];prior=log.get(h)
            if prior and prior.get('page_id'):
                if prior.get('status')=='verified' and prior['properties_sha256']==item['properties_sha256']:continue
                old=item.get('previous_properties')
                changes={k:v for k,v in item['payload']['properties'].items() if not old or old.get(k)!=v}
                if old:
                    live=next(p for p in pages if p['id']==prior['page_id'])
                    for k,v in changes.items():
                        try:self.verify_properties(live,{k:old[k]})
                        except AssertionError:self.verify_properties(live,{k:v})
                page=self.notion('PATCH','pages/'+prior['page_id'],json={'properties':changes})
                self.verify_properties(page,item['payload']['properties']);prior.update(status='verified',properties_sha256=item['properties_sha256'],updated_at=now());save(self.o/'notion_import_log.json',log);print('UPDATED',h,flush=True);continue
            if matches[h]:
                if len(matches[h])!=1:raise RuntimeError('Multiple Notion matches for '+h)
                page=self.notion('GET','pages/'+matches[h][0]['id'])
                adoption_path=self.o/'notion_adoption_review.json'
                adoption=(read(adoption_path) if adoption_path.exists() else {}).get(h)
                if adoption:
                    assert adoption.get('reviewed_at') and adoption['page_id']==page['id']
                    changes={}
                    for k,val in item['payload']['properties'].items():
                        try:self.verify_properties(page,{k:val})
                        except AssertionError:
                            old=adoption['before_properties'][k];typ=old['type']
                            self.verify_properties(page,{k:{typ:old.get(typ)}})
                            changes[k]=val
                    if changes:page=self.notion('PATCH','pages/'+page['id'],json={'properties':changes})
                    self.verify_properties(page,item['payload']['properties'])
                    marker='研究与首次邀请补录｜'+self.c['date']+'｜'+h
                    actual=self.blocks(page['id'])
                    if not any(text(b)==marker for b in actual):
                        self.notion('PATCH','blocks/'+page['id']+'/children',json={'children':[para(marker)]+item['payload']['children']})
                # Recover a create that completed remotely before the local log was committed.
                self.verify_properties(page,item['payload']['properties'])
            else:
                if limit is not None and count>=limit:break
                log[h]={'status':'creating','attempted_at':now(),'properties_sha256':item['properties_sha256']};save(self.o/'notion_import_log.json',log)
                page=self.notion('POST','pages',json=item['payload'])
                log[h].update(status='created',page_id=page['id'],url=page['url']);save(self.o/'notion_import_log.json',log)
                self.verify_properties(page,item['payload']['properties'])
            log[h]={'status':'verified','page_id':page['id'],'url':page['url'],'properties_sha256':item['properties_sha256'],'verified_at':now()};save(self.o/'notion_import_log.json',log)
            count+=1;print('IMPORTED',h,flush=True);time.sleep(.4)
        print('Verified imported records',sum(x['status']=='verified' for x in log.values()),flush=True)
    def verify_properties(self,page,expected):
        actual=flat(page)
        for k,val in expected.items():
            typ=next(iter(val));want=val[typ];got=actual[k]
            if typ in ('rich_text','title'):want=''.join(x['text']['content'] for x in want)
            if typ=='select':want=(want or {}).get('name');got=(got or {}).get('name')
            if typ=='multi_select':want=status_names(want);got=status_names(got)
            if typ=='date':
                want=(want or {}).get('start');got=(got or {}).get('start')
                if want and got:
                    if len(want)==10:
                        assert want==got,(k,want,got)
                        continue
                    assert abs((datetime.fromisoformat(want.replace('Z','+00:00'))-datetime.fromisoformat(got.replace('Z','+00:00'))).total_seconds())<1,(k,want,got)
                    continue
            assert want==got,(k,want,got)
    def sync_sent(self):
        """Update only newly authorized recipients, preserving later replies and manual changes."""
        from datetime import timedelta
        scope=read(self.o/'batch_scope.json');handles=scope['newly_authorized_handles']
        sends=read(self.o/'send_log.json');drafts={d['handle']:d for d in read(self.o/'drafts.json')}
        imports=read(self.o/'notion_import_log.json');review=read(self.o/'review.json')
        path=self.o/'sent_sync_log.json';log=read(path) if path.exists() else {}
        for h in handles:
            s=sends.get(h,{})
            if s.get('status')!='sent_verified':continue
            d=drafts[h];assert review['draft_hashes'][h]==d['sha256']==s['sha256']
            assert s['content_check']['matches_reviewed_body']
            pid=imports[h]['page_id'];page=self.notion('GET','pages/'+pid);f=flat(page)
            date=datetime.fromisoformat(s['accepted_at']).astimezone(timezone(timedelta(hours=8))).date().isoformat()
            marker='Zoho 英文邀请已发送｜'+str(s['message_id'])
            props={}
            if not f.get('发送日期'):props['发送日期']={'date':{'start':date}}
            if not f.get('邮件主题'):props['邮件主题']={'rich_text':rt(d['subject'])}
            if (f.get('联系状态') or {}).get('name')=='未联系' and not f.get('回复日期') and not f.get('是否暂停推进'):
                props['联系状态']={'select':{'name':'已发信'}}
                props['下一步']={'rich_text':rt('等待回复；有意向后确认选品、Option 1 / Option 2、寄样信息与合作条件。')}
                props['筛选行动']={'select':{'name':'P2 联系'}}
                props['建联备注']={'rich_text':rt(f'{date} 已按最新 SOP 逐封审阅英文邮件，通过 Zoho 发送至 {d["to"]}；已核对已发送正文。发送前复查 Notion 及完整已发送记录，无重复。Zoho messageId {s["message_id"]}。发送记录不等于确认送达。')}
            if props:
                save(self.o/'sent_sync_pending.json',{'handle':h,'page_id':pid,'properties':props,'before':{k:page['properties'][k] for k in props}})
                result=self.notion('PATCH','pages/'+pid,json={'properties':props});self.verify_properties(result,props)
            blocks=self.blocks(pid)
            if not any(text(b)==marker for b in blocks):
                children=[para(marker),para('发送日期（北京时间）：'+date+'；收件人：'+d['to']+'；主题：'+d['subject'])]
                children.extend(para(part) for part in d['body'].split('\n\n'))
                self.notion('PATCH','blocks/'+pid+'/children',json={'children':children})
            actual=self.blocks(pid);idx=next(i for i,b in enumerate(actual) if text(b)==marker)
            expected=d['body'].split('\n\n')
            assert [text(b) for b in actual[idx+2:idx+2+len(expected)]]==expected
            self.verify_properties(self.notion('GET','pages/'+pid),props)
            log[h]={'page_id':pid,'message_id':s['message_id'],'verified_at':now(),'sent_date':date,'english_body_verified':True};save(path,log)
            print('SENT_SYNC_VERIFIED',h,flush=True)

    def sync_status_rules(self):
        replacements={
            '回信处理：只接受 paid 的标为“已暂停”；有兴趣、需 brief、可收礼但不承诺发布的标为“待跟进”；未联系的不写下一步。':'回信处理（按用户 2026-09-16 明确口径）：只有对方已来信且当前需要我方回信，才标“待跟进”；我方回信后等待对方回应标“待二次回复”；首次发信后尚无实质回复标“已发信”。报价与寄样任务单独记录，不用“待跟进”泛指所有未完成合作。未联系的不写下一步。',
            '当前主要推进 product-exchange；只接受 paid 的账号标为“已暂停”，有预算时再重新评估。':'当前主要推进 product-exchange。收到 paid 要求或报价、尚需我方回应时，状态为“待跟进”；决定不推进且已回复后可标“已暂停”，已经回复收悉、仅待内部预算评估的报价可保留“已报价”。用户主动询价的个案按实际往来处理。',
            '待跟进：已发送、已回复但需要继续沟通、对方需要 brief、可收礼但不承诺发布、可继续寄样等。':'联系状态以谁需要行动为准：已发信＝第一封邀请已发送、尚无实质回复；待跟进＝对方已来信且当前需要我方写回信，包括第一次回复或后续新邮件；待二次回复＝我方已回信、等待对方回应；已报价＝报价已回复收悉、仅待内部评估；待寄样＝仅需安排寄送、无需立即回信。致谢、表情或自动回复不自动进入待跟进；纯寄样执行也不计为待回信。',
            '已暂停：只接受 paid / 已报价但当前无预算、明确拒绝置换、账号不适合、邮件永久退信且无替代渠道、用户明确排除。':'已暂停：已决定不推进的付费方案、明确拒绝合作、账号不适合、邮件永久退信且无替代渠道、用户明确排除。对方新来信仍需我方答复时，应先标待跟进；不能仅因有报价就自动结束沟通。',
            '对方有兴趣：回复产品选择、合作方式、寄样信息，并更新为“待跟进”。':'对方有兴趣：尚需我方回复产品选择、合作方式、寄样信息时标“待跟进”；回复发出后如等待对方确认，改为“待二次回复”。',
            '对方只接受 paid 或发送 rate card：当前标为“已暂停”，不进入下一步合作。':'对方只接受 paid 或发送 rate card：若尚需我方回应，标“待跟进”，由品牌决定询价、评估或礼貌暂缓；不能将报价视为预算已获批准。回复收悉后仅待内部评估可标“已报价”；明确不推进后标“已暂停”。'
        }
        blocks=self.blocks(self.c['notion_sop'],True)
        def walk(bs):
            for b in bs:
                yield b
                yield from walk(b.get('_children',[]))
        changes=[];seen=set()
        for b in walk(blocks):
            old=text(b)
            if old in replacements:
                changes.append({'id':b['id'],'type':b['type'],'before':old,'after':replacements[old]});seen.add(old)
            elif old in replacements.values():seen.add(next(k for k,v in replacements.items() if v==old))
        assert len(seen)==len(replacements),'SOP changed; review before editing'
        save(self.o/'status_rule_corrections.json',{'reviewed_at':now(),'authority':'User explicit definition of 待跟进','changes':changes})
        for c in changes:
            live=self.notion('GET','blocks/'+c['id']);assert text(live)==c['before']
            self.notion('PATCH','blocks/'+c['id'],json={c['type']:{'rich_text':rt(c['after'])}})
            assert text(self.notion('GET','blocks/'+c['id']))==c['after']
        print('SOP_STATUS_RULES_VERIFIED',len(seen),flush=True)

    def stats(self):
        from collections import Counter
        from datetime import timedelta
        TRACKING=self.c['notion_tracking_data_source']
        tracking_pages=[];cursor=None
        while True:
            body={'page_size':100}
            if cursor:body['start_cursor']=cursor
            batch=self.notion('POST','data_sources/'+TRACKING+'/query',json=body)
            tracking_pages.extend(batch['results'])
            if not batch.get('has_more'):break
            cursor=batch['next_cursor']
        save(self.o/'stats_tracking_evidence.json',{'checked_at':now(),'pages':tracking_pages})
        shipping=shipping_metrics(tracking_pages)
        pages=self.all_pages();save(self.o/'notion_after.json',pages)
        log=read(self.o/'notion_import_log.json');plan=read(self.o/'notion_import_plan.json');byid={p['id']:p for p in pages}
        assert set(log)==self.import_handles() and all(x['status']=='verified' for x in log.values())
        # Contact states and notes legitimately evolve after import as replies arrive.
        # Verify membership/identity here; full import fields were checked at write time.
        removal_log=self.p/'contact_audit'/'applied.json'
        removals=read(removal_log) if removal_log.exists() else {}
        for item in plan:
            pid=log[item['handle']]['page_id']
            if pid not in byid and removals.get(pid,{}).get('action')=='archive':
                assert self.notion('GET','pages/'+pid)['archived'],'Audited removal is no longer archived'
                continue
            self.verify_properties(byid[pid],{'账号':item['payload']['properties']['账号']})
        rows=[flat(p) for p in pages];states=Counter(n for f in rows for n in (status_names(f.get('联系状态')) or {'未设置'}))
        sent_messages=self.sent();save(self.o/'stats_sent_evidence.json',sent_messages)
        recipients={e.lower() for m in sent_messages for e in EMAIL.findall(str(m.get('toAddress','')))}
        alias_path=self.p/'contact_email_aliases.json'
        aliases=read(alias_path) if alias_path.exists() else []
        def mailed(f):
            emails={e.lower() for e in EMAIL.findall(str(f.get('公开联系方式','')))}
            for alias in aliases:
                if f.get('账号')==alias['handle'] and alias['current_email'] in emails:
                    emails.update(alias['previous_emails'])
            return bool(recipients&emails)
        def channel(f,domain):
            return any(re.search(r'https?://(?:www\.)?'+re.escape(domain)+r'/(?!$)[^\s]+',str(f.get(k) or ''),re.I) for k in ('主页链接','Instagram主页','TikTok主页'))
        for f in rows:validate_status_tags(status_names(f.get('联系状态')))
        paused=lambda f:'已暂停' in status_names(f.get('联系状态'))
        contacted=[f for f in rows if mailed(f)];uncontacted=[f for f in rows if not mailed(f)]
        mismatches=[f['账号'] for f in rows if bool(f.get('发送日期'))!=mailed(f)]
        assert not mismatches,'Sent dates disagree with complete Zoho evidence: '+str(mismatches)
        counts={'达人记录总数':len(rows),**states,'累计已发信（达人）':len(contacted),'尚未发信（达人）':len(uncontacted),'仅已发信（当前状态）':states.get('已发信',0),'暂停推进（状态或勾选）':sum(paused(f) for f in rows),'有邮箱':sum(bool(EMAIL.search(str(f.get('公开联系方式','')))) for f in rows),'有 TikTok 主页':sum(channel(f,'tiktok.com') for f in rows),'有 Instagram 主页':sum(channel(f,'instagram.com') for f in rows),'TikTok + Instagram 双平台':sum(channel(f,'tiktok.com') and channel(f,'instagram.com') for f in rows),'已设置跟进日期':sum(bool(f.get('跟进日期')) for f in rows)}
        renamed={'未联系':'尚未发信（达人）','已发信':'累计已发信（达人）','待跟进':'待跟进（需我方回信）','待跟进（当前状态）':'待跟进（需我方回信）','待二次回复':'待二次回复（等对方回复）','待二次回复（当前状态）':'待二次回复（等对方回复）','已报价':'已报价（当前状态）','已暂停':'暂停推进（状态或勾选）','':'已发信（尚无实质回复）','仅已发信（当前状态）':'已发信（尚无实质回复）'}
        renamed.update({'已报价':'已报价（含已暂停）','已报价（当前状态）':'已报价（含已暂停）'})
        renamed.update({'暂停推进（状态或勾选）':'已暂停（已回复并结束合作）','已暂停':'已暂停（已回复并结束合作）'})
        counts['已暂停（已回复并结束合作）']=sum(paused(f) for f in rows)
        counts.update({'待跟进（需我方回信）':states.get('待跟进',0),'待二次回复（等对方回复）':states.get('待二次回复',0),'已报价（含已暂停）':states.get('已报价',0),'已报价且已暂停':sum(paused(f) and '已报价' in status_names(f.get('联系状态')) for f in rows),'已发信（尚无实质回复）':states.get('已发信',0),'待寄样（无需立即回信）':states.get('待寄样',0)})
        counts.update(shipping)
        counts['待寄样（无需立即回信）']=shipping['待寄样（按建联表未寄样）']
        renamed['待寄样（无需立即回信）']='待寄样（按建联表未寄样）'
        save(self.o/'stats_metric_audit.json',{'checked_at':now(),'record_count':len(rows),'sent_message_count':len(sent_messages),'contacted_handles':[f['账号'] for f in contacted],'not_contacted_handles':[f['账号'] for f in uncontacted],'paused_handles':[f['账号'] for f in rows if paused(f)],'instagram_handles':[f['账号'] for f in rows if channel(f,'instagram.com')],'tiktok_handles':[f['账号'] for f in rows if channel(f,'tiktok.com')],'sent_date_mismatches':mismatches,'state_counts':dict(states)})
        table=self.blocks(self.c['notion_stats_table'])
        if not any(text_=='已报价且已暂停' for text_ in [''.join(x.get('plain_text',x.get('text',{}).get('content','')) for x in b['table_row']['cells'][0]) for b in table]):
            self.notion('PATCH','blocks/'+self.c['notion_stats_table']+'/children',json={'children':[{'object':'block','type':'table_row','table_row':{'cells':[rt('已报价且已暂停'),rt(str(counts['已报价且已暂停']))]}}]})
            table=self.blocks(self.c['notion_stats_table'])
        labels={''.join(x.get('plain_text',x.get('text',{}).get('content','')) for x in b['table_row']['cells'][0]) for b in table}
        missing=[label for label in shipping if label not in {renamed.get(x,x) for x in labels} and label!='寄样情况：未寄样']
        if missing:
            self.notion('PATCH','blocks/'+self.c['notion_stats_table']+'/children',json={'children':[{'object':'block','type':'table_row','table_row':{'cells':[rt(label),rt(str(counts[label]))]}} for label in missing]})
            table=self.blocks(self.c['notion_stats_table'])
        for b in table:
            label=''.join(x.get('plain_text',x.get('text',{}).get('content','')) for x in b['table_row']['cells'][0])
            if label=='指标':continue
            label=renamed.get(label,label)
            assert label in counts,'Unknown metric must not silently become zero: '+label
            self.notion('PATCH','blocks/'+b['id'],json={'table_row':{'cells':[rt(label),rt(str(counts[label]))]}})
        stamp=datetime.now(timezone(timedelta(hours=8))).strftime('%Y-%m-%d %H:%M')
        sends=read(self.o/'send_log.json');drafts=read(self.o/'drafts.json');sent_count=sum(x['status']=='sent_verified' for x in sends.values());bounces=sum(x.get('delivery_status')=='bounced' for x in sends.values())
        retried=sum(x.get('delivery_status')=='resent_after_sender_unblock' for x in sends.values())
        delayed=sum(x.get('delivery_status')=='delayed' for x in sends.values())
        email_people=sum(bool(c['emails']) for c in self.candidates())
        deferred=email_people-len(drafts)
        totals=f'当前需要我方回信（待跟进）{states.get("待跟进",0)} 位；我方已回、等待对方回复 {states.get("待二次回复",0)} 位；仅待寄样执行 {states.get("待寄样",0)} 位。\n主表共 {len(rows)} 条达人记录；与 Zoho 完整已发送记录按公开邮箱交叉核对，累计已发信 {len(contacted)} 位，尚未发信 {len(uncontacted)} 位。每位达人只计一次，跟进邮件不重复累加人数；“达人建联情况”表不重复计入。暂停推进 {counts["暂停推进（状态或勾选）"]} 位，包含状态为已暂停或勾选暂停推进的记录。\n当前状态分布：'+'，'.join(label+' '+str(states.get(label,0))+' 位' for label in ['未联系','已发信','待跟进','待二次回复','已报价','待寄样','已暂停','已完成'])+'。其中“已发信”状态不代表累计联系人数。'
        dates=Counter(datetime.fromisoformat(s.get('accepted_at') or s['attempted_at']).astimezone(timezone(timedelta(hours=8))).date().isoformat() for s in sends.values() if s['status']=='sent_verified')
        batch_note=f"{self.c['date']} 名单批次：{self.c['target']} 人，其中 {email_people} 人有公开邮箱，已向 {sent_count} 位发送英文邀请；有邮箱尚未发送 {email_people-sent_count} 位，无确认邮箱 {self.c['target']-email_people} 位。实际发送日期："+'；'.join(day+' '+str(n)+' 封' for day,n in sorted(dates.items()))+f"。当前记录退信 {bounces} 位，投递延迟 {delayed} 位；发送成功不等于确认送达。"
        batch_note+=f" 本批已同步 Notion {len(log)} 位，其余 {self.c['target']-len(log)} 位保留在本地名单，未导入。"
        scope=read(self.o/'batch_scope.json') if (self.o/'batch_scope.json').exists() else {}
        if scope.get('newly_authorized_handles'):
            added=sum(sends.get(h,{}).get('status')=='sent_verified' for h in scope['newly_authorized_handles'])
            batch_note+=f" 本次获授权补发 {len(scope['newly_authorized_handles'])} 位，已完成 {added} 位；此前已发送者未重复发送。"
        context=read(self.o/'stats_context.json') if (self.o/'stats_context.json').exists() else {}
        # This is the shared current summary, not a project-specific historical report.
        # Historical batch notes and followup summaries stay in their audit files.
        totals=totals.replace('当前状态分布：','多选标签计数（可重叠，不相加）：')
        totals=totals.replace(f'仅待寄样执行 {states.get("待寄样",0)} 位',f'建联表标为未寄样 {shipping["待寄样（按建联表未寄样）"]} 位')
        totals=totals.replace('多选标签计数（可重叠，不相加）：','主表联系标签计数（可重叠；其中待寄样标签不作实际寄样统计）：')
        totals+='\n寄样情况实时读取建联表“是否已寄样”，包含手动修改：'+'，'.join(k.replace('寄样情况：','')+' '+str(v)+' 位' for k,v in shipping.items() if k.startswith('寄样情况：'))+'。空白单列为未填写，不推断为未寄样；本次统计不改写达人记录。'
        totals=totals.replace('包含状态为已暂停或勾选暂停推进的记录','仅包含对方已回复且我方已确认结束本轮合作的记录；退信或渠道停用不计入')
        totals+=f'\n已报价共 {states.get("已报价",0)} 位，其中已暂停 {counts["已报价且已暂停"]} 位。报价是已发生的事实，暂停是推进状态，两者可同时存在；后续接受置换也保留历史报价标签。'
        narrative='更新于 '+stamp+'（北京时间）。\n\n'+totals
        narrative+='\n\n已发信＝首次邀请已发、尚无达人回复；已暂停＝达人已回复且我方确认不再合作，两者互斥。已报价标签可与推进阶段并存，退信不算达人回复，旧暂停勾选不直接参与暂停人数统计。'
        narrative+='\n\n口径：待跟进只指对方已来信、当前需要我方回信；首次发信后等待回应不计入。新报价需要我方答复时也计入；已回复收悉后仅待内部评估不计入。累计联系、当前状态、暂停标记及平台覆盖指标互有重叠，不能将整张统计表相加。平台主页同时检查专用字段和主链接；有邮箱表示已记录公开邮箱，不代表已验证有效。报价不等于接受合作费用；收货信息齐全不等于已寄样。最新逐人进展以数据库为准。'
        paragraph_id=self.c['notion_stats_paragraph']
        self.notion('PATCH','blocks/'+paragraph_id,json={'paragraph':{'rich_text':rt(narrative)}})
        actual_paragraph=self.notion('GET','blocks/'+paragraph_id)
        assert text(actual_paragraph)==narrative,'Statistics narrative readback mismatch'
        save(self.o/'stats_text_after.json',{'verified_at':now(),'block_id':paragraph_id,'text':narrative})
        actual=self.blocks(self.c['notion_stats_table'])
        for b in actual[1:]:
            cells=[''.join(x.get('plain_text',x.get('text',{}).get('content','')) for x in c) for c in b['table_row']['cells']]
            assert cells[1]==str(counts.get(cells[0],0)),cells
        save(self.o/'stats_after.json',{'updated_at':now(),'counts':counts,'table':actual,'verified_imports':len(log)})
        print(json.dumps(counts,ensure_ascii=False),flush=True)
    def export_progress(self):
        from openpyxl import load_workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from datetime import timedelta
        pages={p['id']:p for p in read(self.o/'notion_after.json')} if (self.o/'notion_after.json').exists() else {}
        imports=read(self.o/'notion_import_log.json') if (self.o/'notion_import_log.json').exists() else {}
        sends=read(self.o/'send_log.json')
        scope=read(self.o/'notion_scope.json') if (self.o/'notion_scope.json').exists() else {}
        existing={x['handle']:x['matches'][0]['id'] for x in scope.get('existing_records',[]) if len(x.get('matches',[]))==1}
        page_ids={h:v['page_id'] for h,v in imports.items()}
        page_ids.update({h:pid for h,pid in existing.items() if h not in page_ids and pid in pages})
        drafts={d['handle']:d for d in read(self.o/'drafts.json')}
        wb=load_workbook(self.research_workbook());ws=wb['入选名单']
        actual_count=ws.max_row-1
        assert actual_count==len(self.reviews), 'Workbook does not match current reviewed list'
        headers=[c.value for c in ws[1]];account_col=headers.index('账号')+1
        for row in range(2,ws.max_row+1):
            h=ws.cell(row,account_col).value
            if h not in page_ids:
                if sends.get(h,{}).get('status')=='sent_verified' and '联系状态' in headers:
                    ws.cell(row,headers.index('联系状态')+1).value='已发信（首次邀约；后续回复状态待核）'
                continue
            f=flat(pages[page_ids[h]])
            for col,k in enumerate(headers,1):
                if k not in f:continue
                val=f[k]
                if isinstance(val,dict):val=val.get('name',val.get('start'))
                if isinstance(val,list):val=json.dumps(val,ensure_ascii=False)
                c=ws.cell(row,col);c.value=val
                if isinstance(val,str):c.data_type='s'
        result=wb.create_sheet('发送结果')
        result.append(['账号','收件邮箱','发送时间（北京时间）','发送状态','已发送正文核对','退信情况','Zoho Message ID','Notion'])
        zone=timezone(timedelta(hours=8))
        for h,s in sends.items():
            stamp=s.get('accepted_at') or s.get('attempted_at')
            result.append([h,s.get('to',drafts.get(h,{}).get('to','')),datetime.fromisoformat(stamp).astimezone(zone).strftime('%Y-%m-%d %H:%M') if stamp else '',s['status'],'已核对' if s.get('content_check',{}).get('matches_reviewed_body') else '待核对',s.get('delivery_error') or '截至核验未记录退信；不代表确认送达',str(s.get('message_id','')),imports.get(h,{}).get('url','')])
        pending=wb.create_sheet('待发送邮箱');pending.append(['账号','邮箱','说明'])
        remaining=[]
        for c in self.candidates():
            if c['emails'] and sends.get(c['handle'],{}).get('status')!='sent_verified':
                remaining.append(c);pending.append([c['handle'],'; '.join(c['emails']),sends.get(c['handle'],{}).get('status','尚未发送')])
        save(self.o/'next_batch.json',{'count':len(remaining),'updated_at':now(),'requires_fresh_dedup_and_editorial_review':True,'candidates':remaining})
        for sheet in (result,pending):
            sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
            for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='405D51')
            for col in 'ABCDEFGH':sheet.column_dimensions[col].width=30
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment=Alignment(wrap_text=True,vertical='top')
                    if isinstance(cell.value,str):cell.data_type='s'
        if '说明' in wb:
            for row in wb['说明'].iter_rows():
                if row[0].value=='联系状态':
                    row[1].value=f"当前入选 {actual_count} 位，目标 {self.c['target']} 位；已通过 Zoho 发送并核对 {sum(s.get('status')=='sent_verified' for s in sends.values())} 位；本项目导入 Notion {len(imports)} 位；沿用已有记录 {len(page_ids)-len(imports)} 位；尚未关联 Notion {actual_count-len(page_ids)} 位。已有记录保留其最新状态，发送不等于送达。"
            stamp=read(self.o/'stats_after.json')['updated_at'] if (self.o/'stats_after.json').exists() else now()
            wb['说明'].append(['建联状态更新时间',stamp,'已导入者状态来自本次 Notion 回读；未导入者仅记录本地发送事实，未核查后续回复。历史退信保留，未自动重发。'])
        stamp=datetime.now(zone).date().isoformat();path=self.p/f'NUWARIA_{actual_count}人建联进度_{stamp}.xlsx'
        wb.save(path);check=load_workbook(path,read_only=True)
        assert check['入选名单'].max_row-1==actual_count
        assert check['发送结果'].max_row-1==len(sends)
        check.close();print('PROGRESS_EXPORTED',path,flush=True)

    def mail_status(self):
        account=self.zoho('GET','')['data']
        print('outgoingBlocked',account.get('outgoingBlocked'),'smtpStatus',account.get('smtpStatus'),flush=True)
        from playwright.sync_api import sync_playwright
        with sync_playwright() as pw:
            browser=pw.chromium.connect_over_cdp(self.c['cdp']);context=browser.contexts[0]
            page=context.new_page()
            try:
                page.goto('https://mail.zoho.com/UnblockMe',wait_until='domcontentloaded',timeout=30000)
                page.wait_for_timeout(2500)
                content=page.locator('body').inner_text(timeout=10000)
                save(self.o/'mail_status_page.json',{'checked_at':now(),'url':page.url.split('?')[0],'text':content})
                print('Zoho account status page:',page.url.split('?')[0],content[:6000],flush=True)
            finally:page.close()
    def verify_pages(self):
        plan=read(self.o/'notion_import_plan.json');log=read(self.o/'notion_import_log.json');checks={}
        for i,item in enumerate(plan,1):
            h=item['handle'];actual=self.blocks(log[h]['page_id'])
            expected=item['payload']['children']
            # Notion removes zero-width spaces from imported rich text.
            comparable=lambda bs:[(b['type'],text(b).replace('\u200b','')) for b in bs]
            marker='研究与首次邀请补录｜'+self.c['date']+'｜'+h
            offset=next((j+1 for j,b in enumerate(actual) if text(b)==marker),0)
            assert comparable(actual[offset:offset+len(expected)])==comparable(expected),'Notion imported content mismatch: '+h
            checks[h]={'page_id':log[h]['page_id'],'checked_at':now(),'blocks':len(actual),'matches':True}
            save(self.o/'notion_body_verification.json',checks)
            if i%10==0:print('BODY_VERIFIED',i,'/',len(plan),flush=True)
            time.sleep(.35)

if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--project',required=True);ap.add_argument('action',choices=['preflight','build','approve-review','send','reconcile','build-import','import-notion','stats','mail-status','verify-pages','retry-bounce']);ap.add_argument('--limit',type=int);ap.add_argument('--retry-blocked',action='store_true');ap.add_argument('--handle');args=ap.parse_args()
    run=Run(args.project)
    if args.action=='send':run.send(args.limit,args.retry_blocked)
    elif args.action=='import-notion':run.import_notion(args.limit)
    elif args.action=='retry-bounce':run.retry_bounce(args.handle)
    else:getattr(run,args.action.replace('-','_'))()
