"""Evidence validation; aesthetic/brand decisions remain a human/model review."""
import json,re
from collections import defaultdict
from pathlib import Path
from workflow_context import P,DATE,TARGET,CUTOFF,CONFIG
from history_scope import external_history
def load(path):return json.loads(path.read_text(encoding='utf8'))
EMAIL=re.compile(r'[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}')
def validate(reviews=None,complete=False,workbook=None):
 d=reviews if reviews is not None else load(P/'production_reviews.json')
 inc={h:v for h,v in d.items() if v.get('verdict')=='include'};errors=[]
 def need(ok,msg):
  if not ok:errors.append(msg)
 need(len(inc)<=TARGET,'入选人数超过目标')
 if complete:need(len(inc)==TARGET,f'只有{len(inc)}/{TARGET}位')
 excluded=set(load(P/'excluded_handles.json'))
 hist=load(P/'historical_exclusions.json') if (P/'historical_exclusions.json').exists() else {}
 he=set(hist.get('emails',[]));he.update(e.lower() for e in EMAIL.findall(json.dumps(external_history(P,load(P/'existing.json')))))
 hp=set(hist.get('person_keys',[]));seen=set();mail=defaultdict(list);frames=0
 for h,v in inc.items():
  need(not re.search(r'\?{3,}|\ufffd',json.dumps(v,ensure_ascii=False)),h+': 审核备注含乱码或连续问号，须修复后导出')
  need(h.lower() not in excluded,h+': 历史账号重复或已排除')
  identity=v.get('person_key',h).lower();need(identity not in seen and identity not in hp,h+': 人物身份重复');seen.add(identity)
  aliases={a.lower() for a in v.get('aliases',[])}
  need(not aliases&excluded,h+': 历史别名重复')
  try:r=load(P/'profiles'/f'{h}.json');audit=load(P/'post_audits'/h/'audit.json')
  except (OSError,ValueError) as e:errors.append(h+': 缺少可读核验数据 '+str(e));continue
  posts=r.get('posts',[])
  if CONFIG.get('min_followers'):
   count=r.get('user',{}).get('follower_count')
   need(isinstance(count,(int,float)) and count>=CONFIG['min_followers'],h+': 主平台粉丝数未确认或低于下限')
  if CONFIG.get('max_followers'):
   count=r.get('user',{}).get('follower_count')
   need(isinstance(count,(int,float)) and 0<count<=CONFIG['max_followers'],h+': 粉丝数未确认或超过上限')
   for linked in v.get('linked_accounts',[]):
    n=linked.get('followers')
    if isinstance(n,(int,float)):
     need(0<n<=CONFIG['max_followers'],h+': 已知关联平台粉丝数超过上限或无效')
  need(len(posts)>=12 and len(audit)==12,h+': 不是连续12条')
  need([x['date'] for x in posts]==sorted([x['date'] for x in posts],reverse=True),h+': 未按时间排序')
  latest=next((x['date'] for x in posts if x['date']<=DATE),'')
  need(CUTOFF<=latest<=DATE,h+': 最近内容不在项目日期前30天内')
  scenes=v.get('scenes',{});need(len(scenes)>=4 and all(str(k).isdigit() and 1<=int(k)<=12 and bool(t.strip()) for k,t in scenes.items()),h+': 缺少4条具体植入依据')
  for k in ['risks','rationale','product']:need(bool(v.get(k)),h+': 缺少 '+k)
  for i,a in enumerate(audit):
   need(a.get('index')==i+1,h+': 审核序号不一致')
   if i<len(posts):need(a.get('url')==posts[i]['url'],h+': 内容证据链接错位')
   fs=a.get('frames',[]);need(not a.get('error') and len(fs) in (1,4) and not any(f.get('error') for f in fs),h+f': 第{i+1}条抽帧不完整')
   released=CONFIG.get('allow_reviewed_image_deletion') is True and v.get('visual_review_completed') is True
   need(released or (P/'post_audits'/h/f'{i+1:02d}_sheet.jpg').exists(),h+f': 第{i+1}条审核图缺失');frames+=len(fs)
  ee=([v['public_email']] if v['public_email'] else []) if 'public_email' in v else r.get('emails',[])
  if CONFIG.get('require_public_email'):
   need(bool(ee) and all(isinstance(e,str) and EMAIL.fullmatch(e) for e in ee),h+': 本批只收录有已确认公开邮箱的达人')
  for e in ee:mail[e.lower()].append(h);need(e.lower() not in he,h+': 邮箱与历史重复')
 for e,hs in mail.items():need(len(hs)==1,'同一邮箱重复计人: '+','.join(hs))
 if workbook:
  from openpyxl import load_workbook
  w=load_workbook(workbook,read_only=True);rows=list(w['入选名单'].values);headers=rows[0]
  need(list(headers)==list(load(P/'schema.json')['properties']),'工作簿表头与schema不一致')
  need(len(rows)==len(inc)+1,'工作簿人数不一致');need(w['12条内容与植入证据'].max_row==len(inc)*12+1,'证据行数不一致')
  need({r[headers.index('账号')] for r in rows[1:]}=={h.removeprefix('tt__') for h in inc},'工作簿账号不一致')
  sends=load(P/'outreach/send_log.json') if (P/'outreach/send_log.json').exists() else {}
  for row in rows[1:]:
   h=row[headers.index('账号')];sent=sends.get(h,{})
   if sent.get('status')=='sent_verified':
    need(bool(sent.get('content_check',{}).get('matches_reviewed_body')),'发送正文尚未核验：'+h)
    need(bool(row[headers.index('发送日期')]),'缺少已发送日期：'+h)
    need(row[headers.index('邮件主题')]==sent.get('subject'),'已发送主题与日志不一致：'+h)
   else:
    need(row[headers.index('联系状态')]=='未联系','未发送者不应修改联系状态')
    need(all(row[headers.index(k)] in (None,'') for k in ['邮件主题','下一步']),'未发送者应暂停邮件草稿和下一步')
   need(bool(row[headers.index('公开联系方式')]),'缺少公开联系渠道')
  w.close()
 if errors:raise ValueError('\n'.join(errors))
 return {'included':len(inc),'target':TARGET,'unique_people':len(seen),'content_evidence_rows':len(inc)*12,'decoded_frames':frames,'public_email_accounts':len(mail),'history_conflicts':0,'limits':'实际内容为视频抽帧/图文首图；不等于完整观看；受众、转化、合作条件待核。'}
