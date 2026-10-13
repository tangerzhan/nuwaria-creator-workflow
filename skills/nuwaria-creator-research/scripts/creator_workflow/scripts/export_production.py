import json,re,sys,statistics
from pathlib import Path
from datetime import datetime,timezone,timedelta
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
from workflow_context import P, CONFIG, DATE, TARGET, CUTOFF, OUTPUT, CDP
from scoring import DIMENSIONS, summarize
sys.stdout.reconfigure(encoding='utf8')
def load(n):return json.loads((P/n).read_text(encoding='utf8'))
reviews=load('production_reviews.json')
headers=list(load('schema.json')['properties'])
excluded=set(load('excluded_handles.json'))|set(CONFIG.get('excluded_handles',[]))
included={h:v for h,v in reviews.items() if v.get('verdict')=='include'}
score_records=load('scores.json') if (P/'scores.json').exists() else {}
score_summaries={h:summarize(score_records.get(h)) for h in included}
content_threshold=CONFIG.get('content_score_threshold')
if score_records:
    included=dict(sorted(included.items(),key=lambda pair: -(score_summaries[pair[0]]['content'] if score_summaries[pair[0]]['content'] is not None else -1)))
assert not(set(included)&excluded)
wb=Workbook();info=wb.active;info.title='说明'
info.append(['项目',f'{DATE} 新找{TARGET}位达人'])
info.append(['目标人数',TARGET]);info.append(['当前已核验入选',len(included)])
info.append(['完成状态',f'已满{TARGET}位' if len(included)==TARGET else f'尚未满额；还缺{TARGET-len(included)}位。初筛未核实者不计入入选。'])
info.append(['标准','Notion SOP + 用户最新要求；排除健身、户外、宠物、手账、纯穿搭；东方气质不是硬门槛，优先自然生活植入。'])
if (P/'fit_criteria_update.json').exists():
    fit_update=load('fit_criteria_update.json')
    info.append([fit_update['heading'],'\n'.join(fit_update['rules'])])
if CONFIG.get('max_followers'):info.append(['粉丝范围',str(CONFIG.get('min_followers',1))+'–'+str(CONFIG['max_followers'])+'；推荐主平台须达标，关联小号仅作身份记录。不将粉丝少等同于低报价或接受赠礼。'])
info.append(['核验方式','最近连续12条按时间排序；实际视频4个时间点抽帧和公开文案，图文首图；并非完整观看或遍历轮播。评论按需抽查，是否抽查见逐人记录。'])
info.append(['指标限制','仅公开当前可见值；隐藏值留空并标注待后台。普通帖按人工排除明显广告/寄赠/抽奖后计算；中位数与平均播放分开。'])
info.append(['受众限制','所在地为自述，不等于粉丝地区；美国/女性/28岁以上比例待后台。邮箱仅核实公开列示，未验证投递。'])
info.append(['商业复核范围','记录已抓取近期公开内容的实际日期范围；未完整覆盖近90天时明确标注，不宣称排他或全部商单已查清。'])
send_log=load('outreach/send_log.json') if (P/'outreach/send_log.json').exists() else {}
verified_sends={h:s for h,s in send_log.items() if s.get('status')=='sent_verified'}
notion_states={}
if (P/'outreach/notion_before.json').exists():
    for page in load('outreach/notion_before.json'):
        props=page.get('properties',{})
        field=props.get('账号',{})
        handle=''.join(t.get('plain_text',t.get('text',{}).get('content','')) for t in field.get(field.get('type','rich_text'),[])).strip().lstrip('@').lower()
        status=props.get('联系状态',{})
        value=status.get(status.get('type','select')) or []
        values=value if isinstance(value,list) else [value]
        names=list(dict.fromkeys(x['name'] for x in values if isinstance(x,dict) and x.get('name')))
        if handle and names:notion_states[handle]=('；'.join(names),props.get('是否暂停推进',{}).get('checkbox',False))
info.append(['联系状态',f'当前入选名单中本地发送日志已核对{len(set(verified_sends)&set(included))}位；发送不等于送达。主表保留发信前Notion快照中的已有状态，后续来信以Notion及邮箱为准。本次导出不写Notion。' if send_log else '本项目尚无已核验发送记录；名单导出不发送邮件或写Notion。'])
if score_records:
    info.append(['评分口径','保留SOP七维权重。0–5分人工证据评价×权重/5；此数值锚点为本轮执行口径，非声称原SOP已有。缺项待核，完整100分仅七项齐全时给出。内容初评仅自然植入25+品牌适配20，满分45；名单按此排序，不代表完整商业优先级或A/B分层。'])
    info.append(['评分锚点','0=已有明确反证；1=匹配很弱；2=零散、间接；3=多条具体依据但关联较窄或明显限制；4=持续且直接、多场景支持；5=反复主动讲解选物/使用理由且与具体产品材质习惯高度一致。资料缺失不打0分。'])
    if content_threshold is not None:info.append(['内容筛选线',f'{content_threshold}/45及以上进入复核；不折算100分。粉丝、邮箱、去重及主业排除独立执行。初筛依据不足者待核。'])
info.append(['更新时间',datetime.now().isoformat()])
ws=wb.create_sheet('入选名单');ws.append(headers)
ev=wb.create_sheet('12条内容与植入证据');ev.append(['账号','序号','发布日期','实际内容摘要/公开文案短摘','具体植入依据','原生帖子链接','审核图','方法限制'])
metrics=wb.create_sheet('普通帖表现');metrics.append(['账号','样本序号','发布日期','点赞','评论','播放','原生链接','指标口径'])
comments=wb.create_sheet('评论与商单核验');comments.append(['账号','评论核验结论','抽查帖子','可见商单复核范围','商业与冲突','风险/未公开资料'])
contacts=wb.create_sheet('联系渠道证据');contacts.append(['账号','公开邮箱','当前主页','官方外链','读取状态','内容证据摘记'])
channels=wb.create_sheet('平台与粉丝核验');channels.append(['人物主账号','平台','公开主页','粉丝或订阅数','核验日期','核验范围'])
scores=wb.create_sheet('评分与联系排序');scores.append(['内容初评顺序','账号','自然植入 /25','品牌适配 /20','转化 /20','受众 /15','表现 /10','合作 /5','风险质量 /5','内容初评 /45','已评得分','已评权重 /100','完整总分 /100','限制'])
score_evidence=wb.create_sheet('评分依据');score_evidence.append(['账号','维度','满分','原始评价 /5','加权得分','依据或待核原因'])
# Email drafts are paused by the user's latest instruction.
mail=None
seen=set()
for h,v in included.items():
    v=dict(v)
    v.setdefault('business','商业条件待核；未作专项调查。')
    v.setdefault('comments_note','未抽查评论；不作为本轮入选必备条件。')
    v.setdefault('comment_indices',[])
    r=load('profiles/'+h+'.json');u=r['user'];posts=r['posts'];a=load('post_audits/'+h+'/audit.json')
    assert len(posts)>=12 and len(a)==12 and len(v['scenes'])>=4,h
    assert not any(x.get('error') or not x.get('frames') or (len(x['frames'])>1 and sum(not f.get('error') for f in x['frames'])<2) for x in a),h
    frame_failures=sum(bool(f.get('error')) for x in a for f in x.get('frames',[]))
    latest=next(x['date'] for x in posts if x['date']<=DATE)
    assert CUTOFF<=latest<=DATE,(h,latest)
    email=r.get('emails',[])
    if 'public_email' in v:email=[v['public_email']] if v['public_email'] else []
    tt=h.startswith('tt__');platform='TikTok' if tt else 'Instagram';account=r.get('platform_handle',h)
    channels.append([account,platform,r['url'],u.get('follower_count'),r.get('checked_at',DATE),'主平台：12条实际抽帧或图文首图及公开简介；评论核验范围见逐人记录'])
    for linked in v.get('linked_accounts',[]):
        channels.append([account,linked['platform'],linked['url'],linked.get('followers'),linked.get('checked_at',DATE),'官方关联账号；仅核身份和公开粉丝数，未完成该平台内容与评论深查'])
    identity=v.get('person_key',h)
    assert identity not in seen;seen.add(identity)
    ordinary=[(i+1,x) for i,x in enumerate(posts) if i+1 not in v.get('ordinary_exclude',[]) and not x.get('paid')][:10]
    def stat(k):
        vals=[x[k] for _,x in ordinary if isinstance(x.get(k),(float,int))]
        return f'{statistics.median(vals):g}（{len(vals)}/{len(ordinary)}条可见）' if vals else '待达人后台确认（公开不可见）'
    met='普通帖样本'+str(len(ordinary))+'条；点赞中位数 '+stat('likes')+'；评论中位数 '+stat('comments')+'；播放中位数 '+stat('views')
    views=[x['views'] for _,x in ordinary if isinstance(x.get('views'),(float,int))]
    allviews=len(views)==10
    audience='美国受众占比、女性比例及28岁以上比例待达人后台确认。'
    scenes='\n'.join(f'{k}. {s} {posts[int(k)-1]["url"]}' for k,s in v['scenes'].items())
    emails='；'.join(email) if email else v.get('contact_status','公开联系方式待核验')
    d={'是否暂停推进':False,'级别':v.get('level','B'),'平台':platform,'账号':account,'名称':account,'达人姓名':v.get('name',u.get('full_name','')),'主页链接':r['url'],'来源链接':r.get('source',''),'粉丝数':u.get('follower_count'),platform+'粉丝数':u.get('follower_count'),platform+'主页':r['url'],'近期活跃':'最新内容 '+latest+'；核对 '+r.get('checked_at',''),'联系状态':'未联系','筛选行动':'P2终核后联系' if v.get('level','B')=='B' else 'P1终核后联系','公开联系方式':emails,'联系方式来源':r['url'],'建联来源':'原生公开主页与公开合作/互动线索','账号调性分析':v['rationale'],'适配依据':scenes,'商业价值分析':v['business']+'\n'+met,'风险/待核实':v['risks']+'\n'+audience,'创作者所在地':v.get('location','未公开确认'),'推荐产品':v['product'],'合作方式':'拟产品置换，可另议销售佣金；未获达人接受','置信度':'内容与活跃已核；受众/转化/商业条件有限','外部核验备注':'已查看最近连续12条实际抽帧/图文首图；详见证据页。可见内容范围 '+posts[-1]['date']+' 至 '+posts[0]['date']+'，近90天未完整覆盖。\n'+met,'来源表':f'{DATE} 新增{TARGET}位项目','邮件主题':'NUWARIA x '+v.get('name',account)+' Collaboration Invitation','近期内容':'\n'.join(f'{i+1}. {x["date"]} {v.get("topics",{}).get(str(i+1),"见逐条内容证据")} {x["url"]}' for i,x in enumerate(posts[:12]))}
    d['邮件主题']=''
    score=score_summaries[h]
    d['评分']=score['total']
    if score_records:
        d['置信度']+='；评分覆盖权重 '+str(score['coverage'])+'/100'
        d['筛选行动']='内容初评排序后联系；商业条件待核'
        if content_threshold is not None:
            content=score['content']
            outcome='待评分' if content is None else ('达到内容线' if content>=content_threshold else '低于内容线，需重筛')
            d['筛选行动']=f'内容 {content if content is not None else "待核"}/45；阈值 {content_threshold}/45；{outcome}；商业条件待核'
    scores.append([scores.max_row,account,*[row[3] if row[3] is not None else '待核' for row in score['rows']],score['content'],score['subtotal'],score['coverage'],score['total'],'内容排序非完整综合分；缺失不算零分、不折算为100分。原B级只表示内容复核结论。'])
    for row in score['rows']:score_evidence.append([account,*row])
    if v.get('product_fit'):d['账号调性分析']+='\n具体产品关联：'+v['product_fit']
    d['联系方式来源']=v.get('public_email_source',r['url'])
    if allviews:d['平均播放']=round(statistics.mean(views),2)
    if h in verified_sends:
        sent=verified_sends[h]
        stamp=sent.get('accepted_at') or sent.get('attempted_at')
        d['发送日期']=datetime.fromisoformat(stamp).astimezone(timezone(timedelta(hours=8))).date().isoformat() if stamp else ''
        d['联系状态'],d['是否暂停推进']=notion_states.get(h,('已发信',False))
        d['建联备注']='首次邀请已发送并核对正文；后续往来状态以Notion及邮箱为准。'
        d['邮件主题']=sent.get('subject','')
    ws.append([d.get(k,'') for k in headers])
    for i,x in enumerate(posts[:12]):
        excerpt=' '.join(x['caption'].split()[:23])[:140]
        topic=v.get('topics',{}).get(str(i+1),'公开文案短摘：'+excerpt)
        audit=next(t for t in a if t['index']==i+1);fs=audit.get('frames',[])
        ok=sum(not f.get('error') for f in fs)
        method=(f'视频抽帧成功{ok}/{len(fs)}点；非完整观看' if len(fs)>1 else '图文首图；未遍历全部轮播')
        ev.append([account,i+1,x['date'],topic,v['scenes'].get(str(i+1),''),x['url'],str(P/'post_audits'/h/f'{i+1:02d}_sheet.jpg'),method])
    for i,x in ordinary:metrics.append([account,i,x['date'],x.get('likes'),x.get('comments'),x.get('views'),x['url'],'空白=公开不可见，待后台；'+met])
    comments.append([account,v['comments_note'],'\n'.join(posts[i-1]['url'] for i in v['comment_indices']),posts[-1]['date']+' 至 '+posts[0]['date']+'；近90天未完整覆盖',v['business'],v['risks']])
    lf=P/'official_links'/(h+'.json');links=load('official_links/'+h+'.json') if lf.exists() else {}
    pages=links.get('pages',[])
    if not pages:contacts.append([account,emails,r['url'],'当前主页未列外链','已核主页',u.get('biography','')])
    for q in pages:
        excerpt=q.get('text','')[:700]
        if any(mark in excerpt for mark in ('â','ð','Ã','ï¸','\ufffd')) or any(0x80<=ord(c)<=0x9f for c in excerpt):
            excerpt='账号级人工审核摘要：'+v['business']
        contacts.append([account,emails,r['url'],q.get('source') or q.get('url'),q.get('error') or str(q.get('status')),excerpt])
    if email and v.get('public_email_source'):
        contacts.append([account,emails,r['url'],v['public_email_source'],'已核实公开列示','公开邮箱来源；未验证投递。'])
if send_log:
    sent_sheet=wb.create_sheet('发送记录')
    sent_sheet.append(['账号','收件邮箱','发送时间（北京时间）','发送状态','已发送正文核对','Zoho Message ID'])
    for h,s in send_log.items():
        stamp=s.get('accepted_at') or s.get('attempted_at')
        sent_sheet.append([h,s.get('to',''),datetime.fromisoformat(stamp).astimezone(timezone(timedelta(hours=8))).isoformat() if stamp else '',s['status'],'一致' if s.get('content_check',{}).get('matches_reviewed_body') else '未确认',str(s.get('message_id',''))])
for sh in wb:
    sh.freeze_panes='A2';sh.auto_filter.ref=sh.dimensions
    for c in sh[1]:c.font=Font(color='FFFFFF',bold=True);c.fill=PatternFill('solid',fgColor='465747')
    for row in sh.iter_rows(min_row=2):
        for c in row:
            c.alignment=Alignment(vertical='top',wrap_text=True)
            if isinstance(c.value,str):
                if c.value.startswith(('=','+','-','@')):c.data_type='s'
                if c.value.startswith(('https://','http://')) and '\n' not in c.value:c.hyperlink=c.value;c.font=Font(color='0563C1',underline='single')
    for i in range(1,sh.max_column+1):sh.column_dimensions[get_column_letter(i)].width=30
    for row in range(2,sh.max_row+1):sh.row_dimensions[row].height=80
info.column_dimensions['B'].width=115
ev.column_dimensions['D'].width=60;ev.column_dimensions['E'].width=70
out=OUTPUT
wb.save(out)
check=load_workbook(out,read_only=True);assert list(next(check['入选名单'].values))==headers;assert check['入选名单'].max_row-1==len(included)
assert check['评分与联系排序'].max_row-1==len(included)
assert check['评分依据'].max_row-1==len(included)*len(DIMENSIONS)
for row,h in zip(check['评分与联系排序'].iter_rows(min_row=2,values_only=True),included):
    summary=score_summaries[h]
    assert row[12]==summary['total'],(h,'完整总分不可使用局部分数')
    assert row[11]==summary['coverage'],(h,'评分覆盖率不一致')
    assert row[9]==summary['content'],(h,'内容分数不一致')
print(out);print('INCLUDED',len(included),'REMAIN',TARGET-len(included))
