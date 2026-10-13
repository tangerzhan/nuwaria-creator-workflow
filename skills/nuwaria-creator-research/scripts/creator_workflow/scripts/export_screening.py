import json,re,sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from workflow_context import P, CONFIG, DATE, TARGET, CUTOFF, OUTPUT, CDP
sys.stdout.reconfigure(encoding='utf8')
def load(n): return json.loads((P/n).read_text(encoding='utf8'))
from workflow_context import screening
from scoring import summarize
score_records=load('scores.json') if (P/'scores.json').exists() else {}
content_threshold=CONFIG.get('content_score_threshold')
reviews=load('production_reviews.json')
excluded=set(load('excluded_handles.json'))
triage=screening()
wb=Workbook();sh=wb.active;sh.title='初筛工作池'
sh.append(['账号','阶段','初筛结论/理由','主页','公开粉丝数','简介线索','最近发布','近期内容线索（未经逐条核实）','复用资料','来源','自然植入 /25','品牌适配 /20','内容分 /45','内容筛选线 /45','分数阶段','分数结论（不替代硬条件）','评分依据'])
counts={}
records=load('candidate_pool.json') if (P/'candidate_pool.json').exists() else {}
for f in sorted((P/'profiles').glob('*.json')):
 records[f.stem]=json.loads(f.read_text(encoding='utf8'))
for h,r in sorted(records.items()):
 if h in excluded:continue
 v=reviews.get(h,{});t=triage.get(h,{});posts=r.get('posts',[]);u=r.get('user',{})
 state={'include':'内容复核通过｜B备选','exclude':'淘汰','hold':'暂缓'}.get(v.get('verdict')) or {'hold':'暂缓','exclude':'淘汰','pass':'初筛通过｜待核实'}.get(t.get('verdict')) or t.get('stage','待初筛')
 reason=v.get('rationale') or v.get('reason') or t.get('reason','')
 counts[state]=counts.get(state,0)+1
 shot=P/'screenshots'/f'{h}.png'
 reuse=str(shot) if shot.exists() else ('已核实资料见 profiles/'+h+'.json' if (P/'profiles'/f'{h}.json').exists() else '历史摘要见 candidate_pool.json；重新初筛')
 scored=score_records.get(h) or t.get('content_scores',{})
 summary=summarize(scored)
 content=summary['content']
 decision='待评分' if content is None else ('未设置阈值' if content_threshold is None else ('达到内容线' if content>=content_threshold else '低于内容线'))
 stage='内容复核评分' if h in score_records else (t.get('score_stage','概览初评分，待实际内容复核') if scored else '未评分；不默认零分')
 sh.append([h,state,reason,r.get('url'),u.get('follower_count'),u.get('biography',''),posts[0].get('date') if posts else '', '\n'.join(str(i+1)+'. '+' '.join(x.get('caption','').split()[:18])[:120] for i,x in enumerate(posts[:12])),reuse,r.get('source'),*[row[3] if row[3] is not None else '待核' for row in summary['rows'][:2]],content,content_threshold,stage,decision,'\n'.join(row[0]+': '+row[4] for row in summary['rows'][:2])])
sh.freeze_panes='A2';sh.auto_filter.ref=sh.dimensions
for c in sh[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='465747')
for row in sh.iter_rows(min_row=2):
 for c in row:
  c.alignment=Alignment(wrap_text=True,vertical='top')
  if isinstance(c.value,str) and c.value.startswith(('=','+','-','@')):c.data_type='s'
for col in 'ABCDEFGHIJKLMNOPQ':sh.column_dimensions[col].width=28
sh.column_dimensions['F'].width=48;sh.column_dimensions['H'].width=75
info=wb.create_sheet('流程说明');info.append(['阶段','执行内容'])
info.append(['初筛','仅复用简介、近期内容分布、公开画面和明显店铺线索；初筛通过不计入最终100人。'])
info.append(['核实重筛','只对初筛通过者核实实际内容、自然植入、评论、商业冲突与联系渠道。'])
info.append(['邮件','暂停草稿；确定入选后再处理。'])
info.append(['计数',json.dumps(counts,ensure_ascii=False)])
out=P/f'NUWARIA_{DATE}_初筛工作池.xlsx';wb.save(out);print(out);print(counts)
